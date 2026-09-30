#!/usr/bin/env python3
"""Generate private evidence for one API-CONFIG worksheet's contained resources."""

# Copyright ConnectHealth.info (Connecting Solution & Applications Ltd.)
# SPDX-License-Identifier: Apache-2.0

from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys
from tempfile import NamedTemporaryFile

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from adapter_ingestion.ai.terminology import TerminologyCodingAssistant, UnrankedCodingRanker
from adapter_ingestion.manufacturers.registry import get_adapter
from adapter_ingestion.models import AdapterContext
from adapter_ingestion.pipeline import run_pipeline
from adapter_ingestion.service.api_config import extract_embedded_api_config


class EmptyTerminologyClient:
    def search(self, _request):
        return []


def marker_value(marker: object, key: str, fallback: str) -> str:
    prefix = f"{key}="
    for token in str(marker or "").replace(";", ":").split(":"):
        if token.startswith(prefix):
            return token[len(prefix):].strip() or fallback
    return fallback


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--evidence-dir", required=True, type=Path)
    args = parser.parse_args()

    source = load_workbook(args.workbook, read_only=True, data_only=True)
    if args.sheet not in source.sheetnames:
        raise SystemExit(f"worksheet not found: {args.sheet}")
    source_sheet = source[args.sheet]
    marker = source_sheet.cell(1, 1).value
    isolated = Workbook(write_only=True)
    target = isolated.create_sheet(args.sheet)
    for row in source_sheet.iter_rows(values_only=True):
        target.append(list(row))
    source.close()

    with NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
        isolated_path = Path(handle.name)
    isolated.save(isolated_path)
    try:
        embedded = extract_embedded_api_config(isolated_path)
        if embedded is None:
            raise SystemExit("worksheet has no embedded API-CONFIG")
        context = AdapterContext(
            manufacturer="api-config",
            tenant_id="private-local-proof",
            jurisdiction="ES",
            sector="animal-care" if marker_value(marker, "subjectKind", "animal") == "animal" else "health-care",
            issuer_did="did:web:local-proof.invalid",
            audience_did="did:web:local-proof.invalid",
            language=marker_value(marker, "language", "es"),
            subject_kind=marker_value(marker, "subjectKind", "animal"),
            strict_species_mapping=False,
            data_use=marker_value(marker, "dataUse", "secondary"),
            schema_config=embedded["schemaConfig"],
        )
        records = get_adapter("api-config").read_records(isolated_path, context)
        assistant = TerminologyCodingAssistant(
            context=context,
            terminology=EmptyTerminologyClient(),
            ranker=UnrankedCodingRanker(),
        )
        result = run_pipeline(records, context, assistant)
    finally:
        isolated_path.unlink(missing_ok=True)

    examples: dict[str, list[dict]] = {}
    for entry in result.composition_message["body"]["data"]:
        subject = entry["resource"]
        for resource in subject.get("contained", []):
            proposals = resource.get("meta", {}).get("codingProposals", [])
            if not proposals:
                continue
            resource_type = str(resource.get("resourceType", ""))
            bucket = examples.setdefault(resource_type, [])
            if len(bucket) >= 5:
                continue
            bucket.append({
                "subjectId": subject.get("id"),
                "resourceId": resource.get("id"),
                "proposals": proposals,
            })

    document = {
        "sourceWorkbook": str(args.workbook),
        "sheet": args.sheet,
        "records": len(records),
        "summary": result.summary,
        "proposalExamples": examples,
    }
    args.evidence_dir.mkdir(parents=True, exist_ok=True)
    (args.evidence_dir / "contained-resource-review.json").write_text(
        json.dumps(document, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    lines = [
        f"# {args.sheet} contained-resource review proof",
        "",
        f"- Source records: **{len(records)}**",
        f"- Pending coding proposals: **{result.summary['codingProposalEntries']}**",
        "",
        "## Contained resources",
        "",
        "| Resource type | Count |",
        "|---|---:|",
        *(f"| {name} | {count} |" for name, count in result.summary["resourceTypeCounts"].items()),
        "",
        "The JSON contains up to five proposal examples per reviewable resource type.",
    ]
    (args.evidence_dir / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps({"sheet": args.sheet, "summary": result.summary}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
