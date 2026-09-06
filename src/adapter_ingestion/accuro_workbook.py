# Copyright Conéctate Soluciones y Aplicaciones SL
# SPDX-License-Identifier: Apache-2.0

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable
from uuid import UUID, uuid4, uuid5
import re

from gdc_data_utils import ChargeItemClaim, ConditionClaim, InvoiceClaim


ACCURO_MAPPING_REVISION_DATE = "2026-03-19"


@dataclass(frozen=True)
class AccuroSheetConfig:
    name: str
    slug: str
    subject_kind: str
    source_headers: tuple[str, ...]
    internal_fields: tuple[str, ...]
    has_source_header: bool = True
    subject_key_headers: tuple[str, ...] = ()
    output_header_overrides: tuple[tuple[str, str], ...] = ()
    birth_date_headers: tuple[str, ...] = ()
    redacted_headers: tuple[str, ...] = ()
    age_header: str = ""
    birthyear_reference_header: str = ""
    record_date_source_header: str = ""
    invoice_identity_headers: tuple[str, ...] = ()

    @property
    def software_id(self) -> str:
        return f"accuro-{self.slug}"

    @property
    def marker(self) -> str:
        return (
            "API-CONFIG:language=es:"
            f"software-id={self.software_id}:"
            f"subjectKind={self.subject_kind}:dataUse=secondary"
        )

    @property
    def needs_record_date(self) -> bool:
        return "date" not in self.internal_fields

    @property
    def output_headers(self) -> tuple[str, ...]:
        overrides = dict(self.output_header_overrides)
        return tuple(overrides.get(header, header) for header in self.source_headers)

    @property
    def derives_birthyear(self) -> bool:
        return bool(self.age_header and self.birthyear_reference_header)

    @property
    def derives_financial_ids(self) -> bool:
        return bool(self.invoice_identity_headers)

    def synthetic_row(self) -> list[Any]:
        values: list[Any] = []
        for source_header, internal_field in zip(self.source_headers, self.internal_fields):
            if internal_field == "date" or internal_field.endswith("date"):
                values.append(ACCURO_MAPPING_REVISION_DATE)
            elif internal_field == "time":
                values.append("15:00:00")
            elif internal_field == "subject_animal-species":
                values.append("CANINA")
            elif internal_field == "subject_birthsex":
                values.append("Macho")
            elif internal_field == "subject_animal-genderstatus":
                values.append("Esterilizada")
            elif "quantity" in internal_field or internal_field.endswith("items-per-unit"):
                values.append(1)
            else:
                values.append("example")
            if source_header in self.birth_date_headers:
                values[-1] = "2020-05-17"
            elif source_header == self.age_header:
                values[-1] = 5
            elif source_header == self.birthyear_reference_header:
                values[-1] = ACCURO_MAPPING_REVISION_DATE
        return values


ACCURO_SHEET_CONFIGS = (
    AccuroSheetConfig(
        "CV Bestioles", "cv-bestioles", "animal",
        ("MASCOTA", "ESPECIE", "RAZA", "SEXO", "ESTERIL", "EDAD", "ULTIMA VISITA"),
        ("", "subject_animal-species", "subject_animal-breeds", "subject_birthsex", "subject_animal-genderstatus", "", "appointment_lastoccurrencedate"),
        subject_key_headers=("MASCOTA", "ESPECIE", "RAZA", "SEXO"),
        redacted_headers=("MASCOTA",),
        age_header="EDAD",
        birthyear_reference_header="ULTIMA VISITA",
        record_date_source_header="ULTIMA VISITA",
    ),
    AccuroSheetConfig(
        "Canitas 1", "canitas-1", "animal",
        ("EDAD MASCOTA", "MASCOTA", "ESPECIE", "RAZA", "SEXO", "NACIMIENTO"),
        ("", "", "subject_animal-species", "subject_animal-breeds", "subject_birthsex", "subject_birthyear"),
        subject_key_headers=("MASCOTA", "ESPECIE", "RAZA", "SEXO", "NACIMIENTO"),
        output_header_overrides=(("NACIMIENTO", "AÑO NACIMIENTO"),),
        birth_date_headers=("NACIMIENTO",),
        redacted_headers=("MASCOTA",),
    ),
    AccuroSheetConfig(
        "Canitas 2", "canitas-2", "animal",
        ("FECHA", "CONCEPTO", "SECCION", "FAMILIA", "SUBFAMILIA", "DETALLE"),
        ("date", "concept", "section", "family", "subfamily", ""),
        has_source_header=False,
    ),
    AccuroSheetConfig(
        "CV A Caeira", "cv-a-caeira", "animal",
        ("SECCIÓN", "FAMILIA", "SUBFAMILIA", "BASE IMPONIBLE", "IMPORTE IVA", "IMPORTE DESCUENTO", "TOTAL IMPORTE", "CLÍNICA", "TIPO IVA", "EMPRESA NOMBRE FISCAL", "CONCEPTO", "UNIDADES", "CODIGO BARRAS", "IDARTICULO"),
        ("section", "family", "subfamily", "", "", "", "", "", "", "", "concept", "", "", ""),
        redacted_headers=("CLÍNICA", "EMPRESA NOMBRE FISCAL"),
    ),
    AccuroSheetConfig(
        "Crematori de mascotas", "crematori-de-mascotas", "animal",
        ("FECHA", "TIPO DE SERVICIO", "NOMBRE MASCOTA", "ESPECIE", "RAZA", "PESO (KG)", "FECHA NACIMIENTO", "FECHA DEFUNCIÓN", "DETALLE"),
        ("date", "family", "", "subject_animal-species", "subject_animal-breeds", "observation_weight", "subject_birthyear", "subject_deathdate", ""),
        subject_key_headers=("NOMBRE MASCOTA", "FECHA NACIMIENTO", "FECHA DEFUNCIÓN"),
        output_header_overrides=(("FECHA NACIMIENTO", "AÑO NACIMIENTO"),),
        birth_date_headers=("FECHA NACIMIENTO",),
        redacted_headers=("NOMBRE MASCOTA",),
    ),
    AccuroSheetConfig(
        "Mascoverso", "mascoverso", "animal",
        ("nombre", "especie", "edad_anios", "entorno", "ficha", "tos", "estornudos", "vomitos", "rascarse", "jadeo", "quejido", "respiracion", "ladrido", "aullido", "gruñido", "comer", "beber", "total_detecciones", "eventos_collar_totales", "registros_imu", "beacons", "primera_deteccion", "ultima_deteccion", "dias_observado"),
        ("", "subject_animal-species", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "account_serviceperiod-start", "date", ""),
        subject_key_headers=("nombre", "especie"),
        redacted_headers=("nombre",),
    ),
    AccuroSheetConfig(
        "Pinol Vepahi", "pinol-vepahi", "animal",
        ("FechaVisita", "TipoVisita", "ESPECIE", "RAZA", "Anamnesis", "tratamiento", "FECHANACIMIENTO", "EDAD", "CLINICA", "Diagnostico"),
        ("date", "family", "subject_animal-species", "subject_animal-breeds", "concept", "procedure_code-display", "subject_birthyear", "", "", f"coding-input:{ConditionClaim.CODE}"),
        output_header_overrides=(("FECHANACIMIENTO", "AÑO NACIMIENTO"),),
        birth_date_headers=("FECHANACIMIENTO",),
        redacted_headers=("CLINICA",),
    ),
    AccuroSheetConfig(
        "Survet Diagonal", "survet-diagonal", "animal",
        ("DATA VISITA", "TIPUS VISITA", "ESPÈCIE", "SEXE", "DATA NAIXEMENT", "EDAT", "ANAMNESI", "RAÇA", "TRACTAMENT", "peso"),
        ("date", "family", "subject_animal-species", "subject_birthsex", "subject_birthyear", "", "concept", "subject_animal-breeds", "procedure_code-display", "observation_weight"),
        output_header_overrides=(("DATA NAIXEMENT", "ANY NAIXEMENT"),),
        birth_date_headers=("DATA NAIXEMENT",),
    ),
    AccuroSheetConfig(
        "Veterinary Automation 1", "veterinary-automation-1", "animal",
        ("PROVINCIA", "FECHA ALTA", "FECHA DEFUNCION", "SEXO", "ESTERIL", "IGUALA", "ULTIMA VISITA", "ESPECIE", "RAZA", "CARACTER", "EDAD", "MASCOTA"),
        ("location_address-district", "account_serviceperiod-start", "subject_deathdate", "subject_birthsex", "subject_animal-genderstatus", "coverage_status", "appointment_lastoccurrencedate", "subject_animal-species", "subject_animal-breeds", "observation_behavior-assessment", "", ""),
        subject_key_headers=("MASCOTA", "ESPECIE", "RAZA", "SEXO"),
        redacted_headers=("MASCOTA",),
        age_header="EDAD",
        birthyear_reference_header="ULTIMA VISITA",
        record_date_source_header="ULTIMA VISITA",
    ),
    AccuroSheetConfig(
        "Veterinary Automation 2", "veterinary-automation-2", "animal",
        ("FECHA_LINEA", "CONCEPTO_LINEA", "CANTIDAD_LINEA", "FECHA_DOCUMENTO", "CONCEPTO_DOCUMENTO", "CANTIDAD_DOCUMENTO", "FAMILIA", "SUBFAMILIA", "MASCOTA", "ESPECIE", "FECHA DEFUNCION", "DESCRIPCION", "COMUNICACION_IDANIMAL"),
        (ChargeItemClaim.OCCURRENCE, ChargeItemClaim.CODE_TEXT, ChargeItemClaim.QUANTITY_NUMBER, InvoiceClaim.DATE, "", "", "family", "subfamily", "", "subject_animal-species", "subject_deathdate", "procedure_code-display", ""),
        subject_key_headers=("COMUNICACION_IDANIMAL",),
        redacted_headers=("MASCOTA", "COMUNICACION_IDANIMAL"),
        invoice_identity_headers=("COMUNICACION_IDANIMAL", "FECHA_DOCUMENTO"),
    ),
    AccuroSheetConfig(
        "Dr Baron dentistas", "dr-baron-dentistas", "person",
        ("Fecha", "Sexo", "Edad", "Tratamiento", "Patologia Dental", "Num Visitas"),
        ("date", "subject_birthsex", "", "family", f"coding-input:{ConditionClaim.CODE}", ""),
        age_header="Edad",
        birthyear_reference_header="Fecha",
    ),
    AccuroSheetConfig(
        "Sanios", "sanios", "person",
        ("IDENTIFICADOR", "SEXO", "DIRECCIÓN", "EDAD", "PATOLOGÍA", "DETALLE"),
        ("", "subject_birthsex", "", "", f"coding-input:{ConditionClaim.CODE}", ""),
        subject_key_headers=("IDENTIFICADOR",),
        redacted_headers=("IDENTIFICADOR",),
    ),
    AccuroSheetConfig(
        "Centro creciendo", "centro-creciendo", "person",
        ("IDENTIFICADOR", "SEXO", "DIRECCIÓN", "EDAD", "PATOLOGÍA", "Visitas por cliente"),
        ("", "subject_birthsex", "", "", f"coding-input:{ConditionClaim.CODE}", ""),
        subject_key_headers=("IDENTIFICADOR",),
        redacted_headers=("IDENTIFICADOR",),
    ),
)


def _row_values(row: Iterable[Any], width: int) -> list[Any]:
    values = list(row)
    if len(values) < width:
        values.extend([None] * (width - len(values)))
    return values[:width]


def _year_from_value(value: Any) -> int | None:
    if isinstance(value, (datetime, date)):
        return value.year
    if isinstance(value, (int, float)):
        numeric = float(value)
        if 1900 <= numeric <= 2100:
            return int(numeric)
        if numeric > 0:
            return (datetime(1899, 12, 30) + timedelta(days=numeric)).year
        return None
    match = re.search(r"\b(19|20)\d{2}\b", str(value or ""))
    return int(match.group(0)) if match else None


def _derive_birthyear(age_value: Any, reference_value: Any) -> int | None:
    reference_year = _year_from_value(reference_value)
    age_match = re.search(r"\d+(?:[.,]\d+)?", str(age_value or ""))
    if reference_year is None or age_match is None:
        return None
    age = int(float(age_match.group(0).replace(",", ".")))
    if age < 0 or age > 150:
        return None
    return reference_year - age


def _research_subject_uuid(
    namespace: UUID,
    config: AccuroSheetConfig,
    source_row_number: int,
    row: list[Any],
) -> str:
    values_by_header = dict(zip(config.source_headers, row))
    subject_parts = [
        f"{header}={str(values_by_header.get(header) or '').strip()}"
        for header in config.subject_key_headers
        if str(values_by_header.get(header) or "").strip()
    ]
    if subject_parts:
        identity = "|".join(subject_parts)
    else:
        canonical = "|".join(str(value or "").strip() for value in row)
        identity = f"row={source_row_number}|{canonical}"
    return str(uuid5(namespace, f"{config.slug}|{identity}"))


def _financial_identifiers(
    namespace: UUID,
    config: AccuroSheetConfig,
    source_row_number: int,
    values_by_header: dict[str, Any],
) -> tuple[str, str]:
    parts = [
        f"{header}={str(values_by_header.get(header) or '').strip()}"
        for header in config.invoice_identity_headers
    ]
    if not parts or any(part.endswith("=") for part in parts):
        return ("", "")
    invoice_identifier = str(uuid5(namespace, f"{config.slug}|invoice|{'|'.join(parts)}"))
    charge_item_identifier = str(
        uuid5(namespace, f"{config.slug}|charge-item|{invoice_identifier}|row={source_row_number}")
    )
    return (invoice_identifier, charge_item_identifier)


def _load_or_create_namespace(path: Path) -> UUID:
    if path.exists():
        return UUID(path.read_text(encoding="utf-8").strip())
    namespace = uuid4()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(f"{namespace}\n", encoding="utf-8")
    return namespace


def prepare_accuro_workbook(
    source_path: Path,
    output_path: Path,
    *,
    split_dir: Path | None = None,
    namespace_path: Path | None = None,
) -> dict[str, Any]:
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:  # pragma: no cover - optional dependency guard
        raise RuntimeError("Install the Excel extra before preparing Accuro workbooks") from exc

    source = load_workbook(Path(source_path), read_only=True, data_only=False)
    config_by_name = {config.name: config for config in ACCURO_SHEET_CONFIGS}
    unknown_sheets = [name for name in source.sheetnames if name not in config_by_name]
    if unknown_sheets:
        raise ValueError(f"Unsupported Accuro sheets: {', '.join(unknown_sheets)}")

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    resolved_namespace_path = Path(namespace_path) if namespace_path else output_path.with_suffix(".subject-namespace")
    namespace = _load_or_create_namespace(resolved_namespace_path)
    if split_dir is not None:
        split_dir = Path(split_dir)
        split_dir.mkdir(parents=True, exist_ok=True)

    combined = Workbook(write_only=True)
    report: dict[str, Any] = {
        "source": str(source_path),
        "output": str(output_path),
        "subjectNamespaceFile": str(resolved_namespace_path),
        "sheets": {},
    }

    for source_sheet in source.worksheets:
        config = config_by_name[source_sheet.title]
        combined_sheet = combined.create_sheet(config.name)
        split_workbook = Workbook(write_only=True) if split_dir is not None else None
        split_sheet = split_workbook.create_sheet(config.name) if split_workbook is not None else None

        extra_headers = ["RESEARCH_SUBJECT_ID"]
        extra_fields = ["subject_id"]
        if config.needs_record_date:
            extra_headers.append("RECORD_DATE")
            extra_fields.append("date")
        if config.derives_birthyear:
            extra_headers.append("SUBJECT_BIRTHYEAR")
            extra_fields.append("subject_birthyear")
        if config.derives_financial_ids:
            extra_headers.extend(("INVOICE_IDENTIFIER", "CHARGE_ITEM_IDENTIFIER"))
            extra_fields.extend((InvoiceClaim.IDENTIFIER, ChargeItemClaim.IDENTIFIER))
        headers = list(config.output_headers) + extra_headers
        internal_fields = list(config.internal_fields) + extra_fields

        def append(values: list[Any]) -> None:
            combined_sheet.append(values)
            if split_sheet is not None:
                split_sheet.append(values)

        append([config.marker])
        append(internal_fields)
        append(headers)

        rows = source_sheet.iter_rows(values_only=True)
        if config.has_source_header:
            next(rows, None)

        source_rows = 0
        research_subject_ids: set[str] = set()
        for source_row_number, raw_row in enumerate(rows, start=2 if config.has_source_header else 1):
            row = _row_values(raw_row, len(config.source_headers))
            if not any(value not in (None, "") for value in row):
                continue
            source_rows += 1
            subject_uuid = _research_subject_uuid(namespace, config, source_row_number, row)
            research_subject_ids.add(subject_uuid)
            values_by_header = dict(zip(config.source_headers, row))
            output_row = [
                None
                if header in config.redacted_headers
                else _year_from_value(value)
                if header in config.birth_date_headers
                else value
                for header, value in zip(config.source_headers, row)
            ]
            output_row.append(subject_uuid)
            if config.needs_record_date:
                record_date = (
                    values_by_header.get(config.record_date_source_header)
                    if config.record_date_source_header
                    else ACCURO_MAPPING_REVISION_DATE
                )
                output_row.append(record_date)
            if config.derives_birthyear:
                output_row.append(
                    _derive_birthyear(
                        values_by_header.get(config.age_header),
                        values_by_header.get(config.birthyear_reference_header),
                    )
                )
            if config.derives_financial_ids:
                invoice_identifier, charge_item_identifier = _financial_identifiers(
                    namespace,
                    config,
                    source_row_number,
                    values_by_header,
                )
                output_row.extend((invoice_identifier, charge_item_identifier))
            append(output_row)

        report["sheets"][config.name] = {
            "softwareId": config.software_id,
            "sourceRows": source_rows,
            "uniqueResearchSubjects": len(research_subject_ids),
            "recordsSharingResearchSubject": source_rows - len(research_subject_ids),
        }
        if split_workbook is not None and split_dir is not None:
            split_workbook.save(split_dir / f"{config.slug}.xlsx")

    combined.save(output_path)
    source.close()
    return report
