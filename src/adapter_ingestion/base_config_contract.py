"""Explicit classification of the legacy BaseConfig ingestion vocabulary."""

from __future__ import annotations

from dataclasses import dataclass
from enum import Enum

from gdc_data_utils import ChargeItemClaim, CoverageClaim, InvoiceClaim, ProcedureClaim


class BaseConfigFieldKind(str, Enum):
    CONTROL = "control"
    CANONICAL_CLAIM = "canonical-claim"
    DATACONV_EXTENSION = "dataconv-extension"
    PENDING = "pending"


@dataclass(frozen=True)
class BaseConfigField:
    kind: BaseConfigFieldKind
    canonical_claim: str = ""
    note: str = ""


def _control() -> BaseConfigField:
    return BaseConfigField(BaseConfigFieldKind.CONTROL)


def _claim(value: str) -> BaseConfigField:
    return BaseConfigField(BaseConfigFieldKind.CANONICAL_CLAIM, value)


def _extension(note: str = "") -> BaseConfigField:
    return BaseConfigField(BaseConfigFieldKind.DATACONV_EXTENSION, note=note)


def _pending(note: str) -> BaseConfigField:
    return BaseConfigField(BaseConfigFieldKind.PENDING, note=note)


BASE_CONFIG_FIELDS: dict[str, BaseConfigField] = {
    "section": _control(),
    "family": _control(),
    "subfamily": _control(),
    "concept": _control(),
    "date": _control(),
    "time": _control(),
    "origin": _pending("Source role vocabulary is not governed."),
    "personal_id": _control(),
    "relatedperson_id": _control(),
    "subject_id": _control(),
    "subject_address-country": _extension(),
    "subject_address-postalcode": _extension(),
    "subject_animal-species": _extension(),
    "subject_animal-breeds": _extension(),
    "subject_birthyear": _extension(),
    "subject_birthsex": _extension(),
    "subject_deathdate": _extension(),
    "subject_animal-genderstatus": _extension(),
    "subject_gender": _extension(),
    "account_serviceperiod-start": _pending("Account claims are not governed by common-utils."),
    "account_serviceperiod-end": _pending("Account claims are not governed by common-utils."),
    "appointment_lastoccurrencedate": _pending("No canonical last-occurrence claim exists."),
    "encounter_participant-type-display": _pending("Requires coded participant semantics."),
    "encounter_service-type-display": _pending("Requires coded encounter type semantics."),
    "coverage_insurer": _pending("A name is not a safe Coverage.payor reference."),
    "coverage_status": _claim(CoverageClaim.STATUS),
    "coverage_period-start": _claim(CoverageClaim.PERIOD_START),
    "coverage_period-end": _claim(CoverageClaim.PERIOD_END),
    "location_address-postalcode": _extension(),
    "location_address-city": _extension(),
    "location_address-district": _extension(),
    "location_address-state": _extension(),
    "observation_behavior-assessment": _pending("Requires an Observation code and value profile."),
    "observation_behavior-traits": _pending("Requires an Observation code and value profile."),
    "observation_weight": _pending("Requires coded quantity and UCUM unit semantics."),
    "procedure_code-display": _claim(ProcedureClaim.CODE_DISPLAY),
    "procedure_followup-date": _pending("No canonical Procedure search claim exists."),
    "procedure_subpotent-date": _pending("No canonical Procedure search claim exists."),
    "procedure_target-display": _pending("Requires a governed coded target."),
    "invoice_identifier": _claim(InvoiceClaim.IDENTIFIER),
    "invoice_date": _claim(InvoiceClaim.DATE),
    "chargeitem_productcode": _claim(ChargeItemClaim.CODE),
    "chargeitem_productname": _claim(ChargeItemClaim.CODE_TEXT),
    "chargeitem_category": _claim(ChargeItemClaim.CATEGORY),
    "chargeitem_supplier-productcode": _claim(ChargeItemClaim.SUPPLIER_PRODUCT_CODE),
    "chargeitem_quantity": _claim(ChargeItemClaim.QUANTITY_NUMBER),
    "chargeitem_unit": _claim(ChargeItemClaim.QUANTITY_UNIT),
    "chargeitem_items-per-unit": _claim(ChargeItemClaim.ITEMS_PER_UNIT),
    "chargeitem_items-quantity": _claim(ChargeItemClaim.ITEMS_QUANTITY_NUMBER),
    "chargeitem_items-unit": _claim(ChargeItemClaim.ITEMS_QUANTITY_UNIT),
}


def canonical_claim_for_base_config_field(field_name: str) -> str:
    """Return the canonical claim only for fields whose semantics are governed."""

    normalized = str(field_name or "").strip()
    entry = BASE_CONFIG_FIELDS.get(normalized)
    if entry is None or entry.kind is not BaseConfigFieldKind.CANONICAL_CLAIM:
        return ""
    return entry.canonical_claim


def reconcile_base_config_workbook(source_path, output_path) -> None:
    """Copy a mapping workbook and split BaseConfig rows by governance state."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - optional dependency guard
        raise RuntimeError("Install the Excel extra before reconciling BaseConfig") from exc

    workbook = load_workbook(source_path)
    if "BaseConfig" not in workbook.sheetnames:
        raise ValueError("mapping workbook does not contain BaseConfig")
    target_names = {
        BaseConfigFieldKind.CONTROL: "BaseConfig-Control",
        BaseConfigFieldKind.CANONICAL_CLAIM: "BaseConfig-Canonical",
        BaseConfigFieldKind.DATACONV_EXTENSION: "BaseConfig-Extensions",
        BaseConfigFieldKind.PENDING: "BaseConfig-Pending",
    }
    for name in target_names.values():
        if name in workbook.sheetnames:
            del workbook[name]
    targets = {kind: workbook.create_sheet(name) for kind, name in target_names.items()}
    header = [cell.value for cell in workbook["BaseConfig"][1]]
    while header and header[-1] is None:
        header.pop()
    output_header = [*header, "Canonical Claim", "Governance note"]
    for sheet in targets.values():
        sheet.append(output_header)

    for row in workbook["BaseConfig"].iter_rows(min_row=2, values_only=True):
        values = list(row[: len(header)])
        if not values or values[0] is None or not str(values[0]).strip():
            continue
        field_name = str(values[0]).strip()
        values[0] = field_name
        entry = BASE_CONFIG_FIELDS.get(
            field_name,
            _pending("Field is absent from the governed BaseConfig contract."),
        )
        targets[entry.kind].append([*values, entry.canonical_claim, entry.note])

    workbook.save(output_path)
