# Flow contract: enrich every Elysa API-CONFIG sheet with canonical Immunization vaccine coding claims while preserving non-administration rows for human review.
# Copyright ConnectHealth.info (Connecting Solution & Applications Ltd.)
# SPDX-License-Identifier: Apache-2.0

from __future__ import annotations

import csv
from io import StringIO
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from adapter_ingestion.elysa_immunization_coding import enrich_workbook


def _fixture(path: Path) -> None:
    workbook = Workbook()
    canitas = workbook.active
    canitas.title = "Canitas 2"
    canitas.append(["API-CONFIG:language=es:subjectKind=animal"])
    canitas.append(["concept", "subfamily"])
    canitas.append(["CONCEPTO", "SUBFAMILIA"])
    canitas.append(["EURICAN R 10 Dosis (APLICACIÓN DE TRATAMIENTOS)", "Vacunas"])
    canitas.append(["EURICAN SOLVENTE DAP/DAPPi", "Vacunas"])
    canitas.append(["PUREVAX FELV INY 10DSX0,5 ML (APLICACIÓN DE TRATAMIENTOS)", "Vacunas"])

    caeira = workbook.create_sheet("CV A Caeira")
    caeira.append(["API-CONFIG:language=es:subjectKind=animal"])
    caeira.append(["family", "concept"])
    caeira.append(["FAMILIA", "CONCEPTO"])
    caeira.append(["VACUNAS", "VERSICAN PLUS BB ORAL 10 DOSIS"])

    pinol = workbook.create_sheet("Pinol Vepahi")
    pinol.append(["API-CONFIG:language=es:subjectKind=animal"])
    pinol.append(["subject_animal-species", "concept"])
    pinol.append(["ESPECIE", "Anamnesis"])
    pinol.append(["CANINA", "EURICAN R + EURICAN DAPPi + EURICAN L4"])
    pinol.append(["CANINA", "Pendiente de vacunar de rabia"])

    dentists = workbook.create_sheet("Dr Baron dentistas")
    dentists.append(["API-CONFIG:language=es:subjectKind=person"])
    dentists.append(["DiagnosticReport.code-text"])
    dentists.append(["Patologia Dental"])
    dentists.append(["Caries dental"])

    sanios = workbook.create_sheet("Sanios")
    sanios.append(["API-CONFIG:language=es:subjectKind=person"])
    sanios.append(["DiagnosticReport.code-text"])
    sanios.append(["PATOLOGÍA"])
    sanios.append(["ENFERMERÍA-Se administra vacuna gripe"])

    workbook.save(path)


def test_enrichment_adds_canonical_correlated_claims_and_skips_non_administrations(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    target = tmp_path / "coded.xlsx"
    _fixture(source)

    summary = enrich_workbook(source, target)

    workbook = load_workbook(target, data_only=True)
    for sheet in workbook.worksheets:
        mappings = [cell.value for cell in sheet[2]]
        assert mappings[-3:] == [
            "Immunization.vaccine-code",
            "Immunization.vaccine-code-display",
            "Immunization.vaccine-code-text",
        ]

    canitas = workbook["Canitas 2"]
    assert canitas.cell(4, 3).value == "http://www.whocc.no/atcvet|QI07AA02"
    assert canitas.cell(4, 4).value == "rabies virus"
    assert canitas.cell(4, 5).value == "virus de la rabia"
    assert canitas.cell(5, 3).value is None
    assert next(csv.reader(StringIO(canitas.cell(6, 4).value))) == [
        "feline leukaemia, recombinant live canarypox virus"
    ]
    assert workbook["CV A Caeira"].cell(4, 3).value == "http://www.whocc.no/atcvet|QI07AE01"

    pinol = workbook["Pinol Vepahi"]
    assert pinol.cell(4, 3).value == (
        "http://www.whocc.no/atcvet|QI07AA02,"
        "http://www.whocc.no/atcvet|QI07AD04,"
        "http://www.whocc.no/atcvet|QI07AB01"
    )
    assert len(pinol.cell(4, 3).value.split(",")) == len(pinol.cell(4, 4).value.split(","))
    assert len(pinol.cell(4, 3).value.split(",")) == len(pinol.cell(4, 5).value.split(","))
    assert pinol.cell(5, 3).value is None

    assert workbook["Dr Baron dentistas"].cell(4, 2).value is None
    assert workbook["Sanios"].cell(4, 2).value == "http://www.whocc.no/atc|J07BB"
    assert summary.matched_rows == 5
    assert {item.species_codes for item in summary.items} == {"9615", "9685", "9606"}
    workbook.close()


def test_enrichment_is_idempotent_and_does_not_duplicate_claim_columns(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    _fixture(source)

    enrich_workbook(source, first)
    enrich_workbook(first, second)

    workbook = load_workbook(second, read_only=True, data_only=True)
    mappings = [cell.value for cell in workbook["Canitas 2"][2]]
    assert mappings.count("Immunization.vaccine-code") == 1
    assert mappings.count("Immunization.vaccine-code-display") == 1
    assert mappings.count("Immunization.vaccine-code-text") == 1
    workbook.close()
