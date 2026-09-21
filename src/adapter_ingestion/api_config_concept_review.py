"""Row-level review queue derived from embedded API-CONFIG semantics."""

# Copyright ConnectHealth.info (Connecting Solution & Applications Ltd.)
# SPDX-License-Identifier: Apache-2.0

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from .source_concept_classification import classify_source_concept


@dataclass(frozen=True)
class ConceptReviewRow:
    sheet: str
    row_number: int
    resource_type: str
    source_text: str
    confidence: str
    basis: str
    section: str
    family: str
    subfamily: str


@dataclass(frozen=True)
class ConceptReviewResult:
    rows: tuple[ConceptReviewRow, ...]
    counts_by_resource_type: dict[str, int]
    counts_by_sheet: dict[str, int]


def _marker_value(marker: object, key: str) -> str:
    prefix = f"{key}="
    for token in str(marker or "").replace(";", ":").split(":"):
        if token.startswith(prefix):
            return token[len(prefix) :].strip()
    return ""


def classify_workbook(path: Path | str) -> ConceptReviewResult:
    """Read a prepared workbook and return candidates without modifying it."""

    from openpyxl import load_workbook

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    rows: list[ConceptReviewRow] = []
    counts_by_sheet: Counter[str] = Counter()
    try:
        for sheet in workbook.worksheets:
            marker = sheet.cell(1, 1).value
            subject_kind = _marker_value(marker, "subjectKind")
            mappings = [str(cell.value or "").strip() for cell in sheet[2]]
            columns = {mapping: index for index, mapping in enumerate(mappings) if mapping}
            if not any(field in columns for field in ("concept", "treatment")):
                continue
            for row_number, values in enumerate(
                sheet.iter_rows(min_row=4, values_only=True),
                start=4,
            ):
                def value(field: str) -> object:
                    index = columns.get(field)
                    return values[index] if index is not None and index < len(values) else ""

                section = str(value("section") or "").strip()
                family = str(value("family") or "").strip()
                subfamily = str(value("subfamily") or "").strip()
                concept = str(value("concept") or "").strip()
                treatment = str(value("treatment") or "").strip()
                if not concept and not treatment:
                    continue
                for candidate in classify_source_concept(
                    section=section,
                    family=family,
                    subfamily=subfamily,
                    concept=concept,
                    treatment=treatment,
                    subject_kind=subject_kind,
                ):
                    if not candidate.source_text:
                        continue
                    rows.append(ConceptReviewRow(
                        sheet=sheet.title,
                        row_number=row_number,
                        resource_type=candidate.resource_type,
                        source_text=candidate.source_text,
                        confidence=candidate.confidence,
                        basis=candidate.basis,
                        section=section,
                        family=family,
                        subfamily=subfamily,
                    ))
                    counts_by_sheet[sheet.title] += 1
    finally:
        workbook.close()
    counts_by_resource = Counter(item.resource_type for item in rows)
    return ConceptReviewResult(
        rows=tuple(rows),
        counts_by_resource_type=dict(sorted(counts_by_resource.items())),
        counts_by_sheet=dict(sorted(counts_by_sheet.items())),
    )
