# Flow contract: reuse shared test fixtures and canonical types; do not introduce duplicated literals.

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace
from uuid import UUID

import pytest
from openpyxl import Workbook, load_workbook
from gdc_data_utils import ChargeItemClaim, DiagnosticReportClaim, InvoiceClaim

from adapter_ingestion.accuro_workbook import ACCURO_SHEET_CONFIGS, prepare_accuro_workbook
from adapter_ingestion.ai.base import NoopCodingAssistant
from adapter_ingestion.manufacturers.registry import get_adapter
from adapter_ingestion.models import AdapterContext
from adapter_ingestion.pipeline import run_pipeline
from adapter_ingestion.runtime.adapters import InMemorySearchRepository
from adapter_ingestion.service.api_config import extract_embedded_api_config
from adapter_ingestion.service.managers.conversion_search import ConversionSearchManager
from adapter_ingestion.service.research import build_storage_namespace


def _prepared_sheet_for_rows(tmp_path: Path, sheet_config, rows: list[list[object]]):
    source = tmp_path / f"{sheet_config.slug}-source.xlsx"
    prepared = tmp_path / f"{sheet_config.slug}-prepared.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_config.name
    if sheet_config.has_source_header:
        worksheet.append(sheet_config.source_headers)
    for row in rows:
        worksheet.append(row)
    workbook.save(source)
    prepare_accuro_workbook(source, prepared)
    return load_workbook(prepared, read_only=True, data_only=True)[sheet_config.name]


def test_last_appointment_derives_birthyear_and_redacts_animal_name(tmp_path: Path) -> None:
    sheet_config = next(item for item in ACCURO_SHEET_CONFIGS if item.name == "CV Bestioles")
    sheet = _prepared_sheet_for_rows(
        tmp_path,
        sheet_config,
        [["PRIVATE PET NAME", "CANINA", "BREED", "Macho", "Fertil", 13, "2025-12-01"]],
    )
    headers = [cell.value for cell in sheet[3]]
    mappings = [cell.value for cell in sheet[2]]
    values = [cell.value for cell in sheet[4]]

    assert mappings[headers.index("ULTIMA VISITA")] == "appointment_lastoccurrencedate"
    assert values[headers.index("RECORD_DATE")] == "2025-12-01"
    assert values[headers.index("SUBJECT_BIRTHYEAR")] == 2012
    assert values[headers.index("MASCOTA")] is None


def test_external_subject_identifier_is_removed_from_prepared_research_workbook(tmp_path: Path) -> None:
    sheet_config = next(item for item in ACCURO_SHEET_CONFIGS if item.name == "Veterinary Automation 2")
    source_identifier = "external-chip-991"
    row = [
        "2026-01-31T11:52:00", "Consulta", 1, "2026-01-31T11:52:17",
        "Documento", 20, "Servicio", "Consulta", "PRIVATE PET NAME", "CANINA",
        "", "Detalle", source_identifier,
    ]

    sheet = _prepared_sheet_for_rows(tmp_path, sheet_config, [row])
    headers = [cell.value for cell in sheet[3]]
    values = [cell.value for cell in sheet[4]]

    assert values[headers.index("COMUNICACION_IDANIMAL")] is None
    assert source_identifier not in values


def test_full_birthdate_is_replaced_with_year_for_anonymization(tmp_path: Path) -> None:
    sheet_config = next(item for item in ACCURO_SHEET_CONFIGS if item.name == "Canitas 1")
    sheet = _prepared_sheet_for_rows(
        tmp_path,
        sheet_config,
        [[4, "PRIVATE PET NAME", "CANINA", "BREED", "Hembra", "2020-05-17"]],
    )
    headers = [cell.value for cell in sheet[3]]
    mappings = [cell.value for cell in sheet[2]]
    values = [cell.value for cell in sheet[4]]

    birthyear_column = headers.index("AÑO NACIMIENTO")
    assert mappings[birthyear_column] == "subject_birthyear"
    assert values[birthyear_column] == 2020
    assert "2020-05-17" not in values
    assert values[headers.index("MASCOTA")] is None


def test_origin_is_ignored_and_free_text_diagnoses_use_diagnostic_report_code_text() -> None:
    for config in ACCURO_SHEET_CONFIGS:
        assert "origin" not in config.internal_fields

    expected_diagnostic_headers = {
        "Pinol Vepahi": "Diagnostico",
        "Dr Baron dentistas": "Patologia Dental",
        "Sanios": "PATOLOGÍA",
        "Centro creciendo": "PATOLOGÍA",
    }
    for sheet_name, source_header in expected_diagnostic_headers.items():
        config = next(item for item in ACCURO_SHEET_CONFIGS if item.name == sheet_name)
        assert config.internal_fields[config.source_headers.index(source_header)] == DiagnosticReportClaim.CODE_TEXT


def test_financial_accuro_fields_use_canonical_claims_only_when_invoice_identity_is_safe() -> None:
    aggregate_catalog = next(
        item for item in ACCURO_SHEET_CONFIGS if item.name == "CV A Caeira"
    )
    for source_header in ("UNIDADES", "CODIGO BARRAS", "IDARTICULO"):
        assert aggregate_catalog.internal_fields[
            aggregate_catalog.source_headers.index(source_header)
        ] == ""

    invoice_lines = next(
        item for item in ACCURO_SHEET_CONFIGS if item.name == "Veterinary Automation 2"
    )
    assert invoice_lines.internal_fields[
        invoice_lines.source_headers.index("FECHA_DOCUMENTO")
    ] == InvoiceClaim.DATE
    assert invoice_lines.internal_fields[
        invoice_lines.source_headers.index("CONCEPTO_LINEA")
    ] == ChargeItemClaim.CODE_TEXT
    assert invoice_lines.internal_fields[
        invoice_lines.source_headers.index("CANTIDAD_LINEA")
    ] == ChargeItemClaim.QUANTITY_NUMBER


def test_invoice_lines_derive_one_stable_invoice_and_distinct_charge_item_ids(tmp_path: Path) -> None:
    config = next(
        item for item in ACCURO_SHEET_CONFIGS if item.name == "Veterinary Automation 2"
    )
    base = [
        "2026-01-31T11:52:00",
        "Consulta",
        1,
        "2026-01-31T11:52:17",
        "Documento",
        2,
        "consultas",
        "seguimiento",
        "PRIVATE PET NAME",
        "CANINA",
        None,
        "Consulta seguimiento",
        854594,
    ]
    second = list(base)
    second[1] = "Medicamento"
    sheet = _prepared_sheet_for_rows(tmp_path, config, [base, second])
    headers = [cell.value for cell in sheet[3]]
    mappings = [cell.value for cell in sheet[2]]
    first = [cell.value for cell in sheet[4]]
    second_output = [cell.value for cell in sheet[5]]

    invoice_column = headers.index("INVOICE_IDENTIFIER")
    charge_column = headers.index("CHARGE_ITEM_IDENTIFIER")
    assert mappings[invoice_column] == InvoiceClaim.IDENTIFIER
    assert mappings[charge_column] == ChargeItemClaim.IDENTIFIER
    assert first[invoice_column] == second_output[invoice_column]
    assert first[charge_column] != second_output[charge_column]
    UUID(first[invoice_column])
    UUID(first[charge_column])


def test_repeated_source_subject_reuses_one_research_subject_uuid(tmp_path: Path) -> None:
    sheet_config = next(item for item in ACCURO_SHEET_CONFIGS if item.name == "CV Bestioles")
    source = tmp_path / "repeated-subject.xlsx"
    prepared = tmp_path / "repeated-subject-api-config.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_config.name
    worksheet.append(sheet_config.source_headers)
    first_row = sheet_config.synthetic_row()
    second_row = list(first_row)
    second_row[sheet_config.source_headers.index("ULTIMA VISITA")] = "2026-03-20"
    worksheet.append(first_row)
    worksheet.append(second_row)
    workbook.save(source)

    prepare_accuro_workbook(source, prepared)

    output = load_workbook(prepared, read_only=True, data_only=True)[sheet_config.name]
    subject_column = [cell.value for cell in output[3]].index("RESEARCH_SUBJECT_ID") + 1
    assert output.cell(4, subject_column).value == output.cell(5, subject_column).value


@pytest.mark.parametrize("sheet_config", ACCURO_SHEET_CONFIGS, ids=lambda item: item.slug)
def test_each_accuro_sheet_gets_api_config_and_imports_without_duplicate_subjects(
    tmp_path: Path,
    sheet_config,
) -> None:
    source = tmp_path / "accuro-source.xlsx"
    prepared = tmp_path / "accuro-api-config.xlsx"
    split_dir = tmp_path / "organizations"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_config.name
    if sheet_config.has_source_header:
        worksheet.append(sheet_config.source_headers)
    worksheet.append(sheet_config.synthetic_row())
    workbook.save(source)

    report = prepare_accuro_workbook(source, prepared, split_dir=split_dir)

    assert report["sheets"][sheet_config.name]["sourceRows"] == 1
    organization_workbook = split_dir / f"{sheet_config.slug}.xlsx"
    extracted = extract_embedded_api_config(organization_workbook)
    assert extracted is not None
    assert extracted["runtimeDefaults"]["softwareId"] == sheet_config.software_id
    assert extracted["runtimeDefaults"]["dataUse"] == "secondary"
    assert extracted["schemaConfig"]["fieldMap"]["subject_id"] == "RESEARCH_SUBJECT_ID"

    adapter = get_adapter("api-config")
    records = adapter.read_records(
        organization_workbook,
        AdapterContext(
            manufacturer=sheet_config.software_id,
            tenant_id=sheet_config.slug,
            jurisdiction="ES",
            sector="onehealth-research",
            issuer_did="did:web:issuer.example",
            audience_did="did:web:audience.example",
            subject_did_prefix="urn:uuid",
            subject_kind=sheet_config.subject_kind,
            strict_species_mapping=False,
            schema_config=extracted["schemaConfig"],
            data_use="secondary",
        ),
    )
    assert len(records) == 1
    subject_identifier = records[0].subject_id
    assert subject_identifier.startswith("urn:uuid:")
    UUID(subject_identifier.removeprefix("urn:uuid:"))

    rerun_split_dir = tmp_path / "organizations-rerun"
    prepare_accuro_workbook(source, prepared, split_dir=rerun_split_dir)
    rerun_records = adapter.read_records(
        rerun_split_dir / f"{sheet_config.slug}.xlsx",
        AdapterContext(
            manufacturer=sheet_config.software_id,
            tenant_id=sheet_config.slug,
            jurisdiction="ES",
            sector="onehealth-research",
            issuer_did="did:web:issuer.example",
            audience_did="did:web:audience.example",
            subject_did_prefix="urn:uuid",
            subject_kind=sheet_config.subject_kind,
            strict_species_mapping=False,
            schema_config=extracted["schemaConfig"],
            data_use="secondary",
        ),
    )
    assert rerun_records[0].subject_id == subject_identifier

    independent_output = tmp_path / "independent-api-config.xlsx"
    independent_split_dir = tmp_path / "independent-organizations"
    prepare_accuro_workbook(source, independent_output, split_dir=independent_split_dir)
    independent_config = extract_embedded_api_config(
        independent_split_dir / f"{sheet_config.slug}.xlsx"
    )
    assert independent_config is not None
    independent_records = adapter.read_records(
        independent_split_dir / f"{sheet_config.slug}.xlsx",
        AdapterContext(
            manufacturer=sheet_config.software_id,
            tenant_id=sheet_config.slug,
            jurisdiction="ES",
            sector="onehealth-research",
            issuer_did="did:web:issuer.example",
            audience_did="did:web:audience.example",
            subject_did_prefix="urn:uuid",
            subject_kind=sheet_config.subject_kind,
            strict_species_mapping=False,
            schema_config=independent_config["schemaConfig"],
            data_use="secondary",
        ),
    )
    assert independent_records[0].subject_id != subject_identifier

    pipeline_result = run_pipeline(records, records_context := AdapterContext(
        manufacturer=sheet_config.software_id,
        tenant_id=sheet_config.slug,
        jurisdiction="ES",
        sector="onehealth-research",
        issuer_did="did:web:issuer.example",
        audience_did="did:web:audience.example",
        subject_did_prefix="urn:uuid",
        subject_kind=sheet_config.subject_kind,
        strict_species_mapping=False,
        schema_config=extracted["schemaConfig"],
        data_use="secondary",
    ), NoopCodingAssistant())
    research_subject = pipeline_result.composition_message["body"]["data"][0]["resource"]
    assert research_subject["resourceType"] == "ResearchSubject"

    repository = InMemorySearchRepository()
    vault_id = build_storage_namespace(
        network_kind="test",
        jurisdiction=records_context.jurisdiction,
        sector=records_context.sector,
        tenant_id=records_context.tenant_id,
    )
    repository.upsert(vault_id=vault_id, resource_type="ResearchSubject", resource=research_subject)
    search_result = ConversionSearchManager(
        SimpleNamespace(settings=SimpleNamespace(demo_mode=True, network_mode="test"), search_repo=repository)
    ).handle(
        tenant_id=records_context.tenant_id,
        jurisdiction=records_context.jurisdiction,
        sector=records_context.sector,
        resource_type="ResearchSubject",
        response=SimpleNamespace(),
        request=SimpleNamespace(headers={}, query_params={}),
        body={
            "resourceType": "Parameters",
            "parameter": [{"name": "identifier", "valueUri": subject_identifier}],
        },
    )
    assert search_result["resourceType"] == "Bundle"
    assert search_result["type"] == "searchset"
    assert search_result["total"] == 1

    prepared_workbook = load_workbook(prepared, read_only=True, data_only=True)
    prepared_sheet = prepared_workbook[sheet_config.name]
    assert str(prepared_sheet.cell(1, 1).value).startswith("API-CONFIG:")
    assert "subject_id" in [cell.value for cell in prepared_sheet[2]]
    assert "RESEARCH_SUBJECT_ID" in [cell.value for cell in prepared_sheet[3]]
