# Flow contract: source workbooks end at concept; confidential identity resolution and derived FHIR resources remain internal.

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from adapter_ingestion.ai.base import NoopCodingAssistant
from adapter_ingestion.manufacturers.registry import get_adapter
from adapter_ingestion.models import AdapterContext
from adapter_ingestion.pipeline import run_pipeline
from adapter_ingestion.research_search_workbooks import (
    FIXTURE_SUBJECT_UUID_BY_ORIGINAL_ID,
    HUMAN_SUBJECTS,
    RESEARCH_STUDY_A,
    RESEARCH_STUDY_B,
    VETERINARY_SUBJECTS,
    generate_research_search_workbooks,
)
from adapter_ingestion.runtime.adapters import InMemoryVaultRepository
from adapter_ingestion.runtime import PreconversionControlPlane
from adapter_ingestion.runtime.adapters import InMemoryBlobStore, InMemoryConfigStore, InMemoryJobQueue, InMemoryJobStore
from adapter_ingestion.runtime.adapters.search import InMemorySearchRepository
from adapter_ingestion.service.api_config import extract_embedded_api_config
from adapter_ingestion.service.managers.conversion_search import ConversionSearchManager
from adapter_ingestion.service.research import build_storage_namespace


EXPECTED_FILES = {
    "veterinary-study-a": "research-search-veterinary-study-a.xlsx",
    "veterinary-study-b": "research-search-veterinary-study-b.xlsx",
    "human-study-a": "research-search-human-study-a.xlsx",
    "human-study-b": "research-search-human-study-b.xlsx",
}
SOURCE_HEADERS = [
    "FECHA",
    "ORIGINAL_PATIENT_ID",
    "AÑO_NACIMIENTO",
    "ESPECIE",
    "SEXO",
    "SECCION",
    "FAMILIA",
    "SUBFAMILIA",
    "CONCEPTO",
]


def _sheet_values(path: Path, sheet_name: str) -> list[tuple[object, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    values = list(workbook[sheet_name].iter_rows(values_only=True))
    workbook.close()
    return values


def _data_rows(path: Path) -> tuple[list[str], list[tuple[object, ...]]]:
    values = _sheet_values(path, "DATA")
    return [str(value or "") for value in values[2]], values[3:]


def _search_cases(path: Path) -> list[dict[str, object]]:
    values = _sheet_values(path, "SEARCH_CASES")
    headers = [str(value or "") for value in values[0]]
    return [dict(zip(headers, row)) for row in values[1:]]


def _workbook_values(path: Path) -> dict[str, list[tuple[object, ...]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    values = {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}
    workbook.close()
    return values


def _readme_value(path: Path, key: str) -> str:
    return {
        str(row[0] or ""): str(row[1] or "")
        for row in _sheet_values(path, "README")[1:]
    }[key]


def _pipeline_subjects(path: Path) -> list[dict[str, object]]:
    embedded = extract_embedded_api_config(path)
    assert embedded is not None
    context = AdapterContext(
        manufacturer="api-config",
        tenant_id="research-fixture",
        jurisdiction="ES",
        sector="onehealth-research",
        issuer_did="did:web:issuer.example",
        audience_did="did:web:audience.example",
        subject_did_prefix="urn:uuid",
        subject_kind=embedded["runtimeDefaults"]["subjectKind"],
        strict_species_mapping=False,
        data_use="secondary",
        personal_id_resolver=lambda original_id: FIXTURE_SUBJECT_UUID_BY_ORIGINAL_ID[str(original_id)],
        schema_config=embedded["schemaConfig"],
    )
    records = get_adapter("api-config").read_records(path, context)
    result = run_pipeline(records, context, NoopCodingAssistant())
    subjects = [entry["resource"] for entry in result.composition_message["body"]["data"]]
    study = _readme_value(path, "RESEARCH_STUDY")
    for subject in subjects:
        subject["meta"]["claims"]["ResearchSubject.study"] = study
    return subjects


def _matches_claim_filter(actual: str, expected: str) -> bool:
    for operator in ("ge", "le", "gt", "lt"):
        if expected.startswith(operator):
            target = expected[2:]
            return {
                "ge": actual >= target,
                "le": actual <= target,
                "gt": actual > target,
                "lt": actual < target,
            }[operator]
    return actual == expected


def _execute_fixture_plan(
    subjects: list[dict[str, object]],
    plan: list[dict[str, object]],
) -> set[str]:
    repository = InMemorySearchRepository()
    for subject in subjects:
        repository.upsert(vault_id="fixture", resource_type="ResearchSubject", resource=subject)
        for resource in subject["contained"]:
            repository.upsert(
                vault_id="fixture",
                resource_type=str(resource["resourceType"]),
                resource=resource,
            )

    candidates = {
        str(subject["meta"]["claims"]["ResearchSubject.identifier"])
        for subject in subjects
    }
    for step in plan:
        resource_type = str(step["resourceType"])
        if resource_type == "ResearchSubject":
            claim_filters = step.get("claimFilters", {})
            assert isinstance(claim_filters, dict)
            matching = {
                str(subject["meta"]["claims"]["ResearchSubject.identifier"])
                for subject in subjects
                if all(
                    _matches_claim_filter(
                        str(subject["meta"]["claims"].get(claim_name, "")),
                        str(expected),
                    )
                    for claim_name, expected in claim_filters.items()
                )
            }
        else:
            params = step["params"]
            assert isinstance(params, dict)
            resources = repository.search(
                vault_id="fixture",
                resource_type=resource_type,
                search_params=params,
            )
            matching = {
                str(resource["meta"]["claims"][f"{resource_type}.subject"])
                for resource in resources
            }
        candidates &= matching
    return candidates


def test_source_workbooks_end_at_concept_and_resolve_original_ids_privately(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    assert {key: path.name for key, path in generated.items()} == EXPECTED_FILES

    for key, path in generated.items():
        workbook = load_workbook(path, read_only=True, data_only=True)
        assert workbook.sheetnames == ["DATA", "SEARCH_CASES", "README"]
        mappings = [str(cell.value or "") for cell in workbook["DATA"][2]]
        workbook.close()

        headers, rows = _data_rows(path)
        assert headers == SOURCE_HEADERS
        assert "RESEARCH_SUBJECT_ID" not in headers
        assert mappings == [
            "date",
            "personal_id",
            "subject_birthyear",
            "subject_animal-species",
            "subject_birthsex",
            "section",
            "family",
            "subfamily",
            "concept",
        ]
        original_ids = [str(row[1]) for row in rows]
        assert all(value.startswith("SYNTHETIC-") for value in original_ids)
        if key.endswith("study-a"):
            counts = Counter(original_ids)
            assert len(counts) == 3
            assert max(counts.values()) >= 3
        else:
            assert len(set(original_ids)) == 1


def test_same_original_id_produces_same_research_subject_across_imports(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    first = _pipeline_subjects(generated["human-study-a"])
    later = _pipeline_subjects(generated["human-study-a"])
    first_ids = {subject["ResearchSubject.identifier"] for subject in first}
    later_ids = {subject["ResearchSubject.identifier"] for subject in later}

    assert first_ids == later_ids
    assert first_ids == {
        f"urn:uuid:{FIXTURE_SUBJECT_UUID_BY_ORIGINAL_ID[f'SYNTHETIC-HUMAN-PATIENT-{index:03d}']}"
        for index in (1, 2, 3)
    }


def test_human_study_matrix_contains_positive_negative_and_isolation_controls(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    headers_a, rows_a = _data_rows(generated["human-study-a"])
    _, rows_b = _data_rows(generated["human-study-b"])
    identifier_index = headers_a.index("ORIGINAL_PATIENT_ID")
    concept_index = headers_a.index("CONCEPTO")

    concepts_a: dict[str, set[str]] = {}
    for row in rows_a:
        concepts_a.setdefault(str(row[identifier_index]), set()).add(str(row[concept_index]))
    concepts_d = {str(row[concept_index]) for row in rows_b}

    subject_a = concepts_a["SYNTHETIC-HUMAN-PATIENT-001"]
    subject_b = concepts_a["SYNTHETIC-HUMAN-PATIENT-002"]
    subject_c = concepts_a["SYNTHETIC-HUMAN-PATIENT-003"]
    assert {"DIABETES MELLITUS", "HbA1c 8,4 %", "METFORMINA 850 MG"} <= subject_a
    assert {"DIABETES MELLITUS", "HbA1c 6,2 %"} <= subject_b
    assert "HbA1c 9,1 %" in subject_c
    assert "DIABETES MELLITUS" not in subject_c
    assert subject_a == concepts_d
    assert _readme_value(generated["human-study-a"], "RESEARCH_STUDY") == RESEARCH_STUDY_A
    assert _readme_value(generated["human-study-b"], "RESEARCH_STUDY") == RESEARCH_STUDY_B


def test_clean_source_rows_materialize_expected_resource_families(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    for key, path in generated.items():
        subjects = _pipeline_subjects(path)
        resource_types = {
            str(resource.get("resourceType"))
            for subject in subjects
            for resource in subject["contained"]
        }
        if key.startswith("human"):
            assert {"Condition", "Immunization", "MedicationStatement", "Observation"} <= resource_types
        else:
            assert {
                "AllergyIntolerance",
                "Condition",
                "DiagnosticReport",
                "Encounter",
                "Immunization",
                "MedicationStatement",
                "Procedure",
            } <= resource_types


def test_every_search_plan_returns_expected_subjects_inside_its_study(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    for path in generated.values():
        subjects = _pipeline_subjects(path)
        study = _readme_value(path, "RESEARCH_STUDY")
        for case in _search_cases(path):
            assert case["RESEARCH_STUDY"] == study
            plan = json.loads(str(case["FHIR_SEARCH_PLAN_JSON"]))
            expected = set(json.loads(str(case["EXPECTED_RESEARCH_SUBJECT_IDS_JSON"])))
            assert _execute_fixture_plan(subjects, plan) == expected, case["CASE_ID"]


def test_identical_human_profiles_remain_isolated_by_research_study(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    study_a = _pipeline_subjects(generated["human-study-a"])
    study_b = _pipeline_subjects(generated["human-study-b"])
    case_a = next(
        case for case in _search_cases(generated["human-study-a"])
        if case["CASE_ID"] == "HUM-C-DIABETES-HBA1C"
    )
    case_b = next(
        case for case in _search_cases(generated["human-study-b"])
        if case["CASE_ID"] == "HUM-C-DIABETES-HBA1C"
    )
    plan = json.loads(str(case_a["FHIR_SEARCH_PLAN_JSON"]))
    matches_a = _execute_fixture_plan(study_a, plan)
    matches_b = _execute_fixture_plan(study_b, plan)

    assert matches_a == set(json.loads(str(case_a["EXPECTED_RESEARCH_SUBJECT_IDS_JSON"])))
    assert matches_b == set(json.loads(str(case_b["EXPECTED_RESEARCH_SUBJECT_IDS_JSON"])))
    assert matches_a.isdisjoint(matches_b)


def test_research_subject_parameters_intersect_resource_collections_inside_one_study(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    repository = InMemorySearchRepository()
    vault_id = build_storage_namespace(
        network_kind="test", jurisdiction="ES", sector="onehealth-research", tenant_id="fixture",
    )
    for workbook_key in ("human-study-a", "human-study-b"):
        for subject in _pipeline_subjects(generated[workbook_key]):
            repository.upsert(vault_id=vault_id, resource_type="ResearchSubject", resource=subject)
            for resource in subject["contained"]:
                repository.upsert(
                    vault_id=vault_id,
                    resource_type=str(resource["resourceType"]),
                    resource=resource,
                )

    manager = ConversionSearchManager(SimpleNamespace(
        settings=SimpleNamespace(demo_mode=True, network_mode="test"),
        search_repo=repository,
    ))
    result = manager.handle(
        tenant_id="fixture",
        jurisdiction="ES",
        sector="onehealth-research",
        resource_type="ResearchSubject",
        response=SimpleNamespace(),
        request=SimpleNamespace(headers={}, query_params={}),
        body={
            "resourceType": "Parameters",
            "parameter": [
                {"name": "ResearchSubject.study", "valueReference": {"reference": RESEARCH_STUDY_A}},
                {"name": "Condition.code:text", "valueString": "diabetes"},
                {"name": "Observation.code:text", "valueString": "hba1c"},
                {"name": "Observation.value-quantity", "valueString": "gt8"},
                {"name": "Observation.date", "valueString": "ge2025-01-01"},
                {"name": "MedicationStatement.code:text", "valueString": "metformina"},
            ],
        },
    )

    assert result["resourceType"] == "Bundle"
    assert result["type"] == "searchset"
    assert {
        entry["resource"]["meta"]["claims"]["ResearchSubject.identifier"]
        for entry in result["entry"]
    } == {HUMAN_SUBJECTS[0]}


def test_every_excel_search_case_executes_as_one_research_subject_parameters_request(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    repository = InMemorySearchRepository()
    vault_id = build_storage_namespace(
        network_kind="test", jurisdiction="ES", sector="onehealth-research", tenant_id="fixture",
    )
    subjects_by_workbook: dict[str, list[dict[str, object]]] = {}
    for workbook_key, path in generated.items():
        subjects_by_workbook[workbook_key] = _pipeline_subjects(path)
        for subject in subjects_by_workbook[workbook_key]:
            repository.upsert(vault_id=vault_id, resource_type="ResearchSubject", resource=subject)
            for resource in subject["contained"]:
                repository.upsert(
                    vault_id=vault_id,
                    resource_type=str(resource["resourceType"]),
                    resource=resource,
                )

    manager = ConversionSearchManager(SimpleNamespace(
        settings=SimpleNamespace(demo_mode=True, network_mode="test"),
        search_repo=repository,
    ))
    for workbook_key, path in generated.items():
        for case in _search_cases(path):
            parameters: list[dict[str, object]] = [{
                "name": "ResearchSubject.study",
                "valueReference": {"reference": str(case["RESEARCH_STUDY"])},
            }]
            for step in json.loads(str(case["FHIR_SEARCH_PLAN_JSON"])):
                resource_type = str(step["resourceType"])
                for parameter, value in step["params"].items():
                    if resource_type == "ResearchSubject" and parameter == "study":
                        continue
                    parameters.append({"name": f"{resource_type}.{parameter}", "valueString": str(value)})
                for claim, value in step.get("claimFilters", {}).items():
                    public_parameter = "birthyear" if claim == "Subject.birthyear" else str(claim)
                    parameters.append({"name": f"ResearchSubject.{public_parameter}", "valueString": str(value)})

            result = manager.handle(
                tenant_id="fixture",
                jurisdiction="ES",
                sector="onehealth-research",
                resource_type="ResearchSubject",
                response=SimpleNamespace(),
                request=SimpleNamespace(headers={}, query_params={}),
                body={"resourceType": "Parameters", "parameter": parameters},
            )
            actual = {
                entry["resource"]["meta"]["claims"]["ResearchSubject.identifier"]
                for entry in result["entry"]
            }
            assert actual == set(json.loads(str(case["EXPECTED_RESEARCH_SUBJECT_IDS_JSON"]))), (
                workbook_key, case["CASE_ID"], actual,
            )


def test_exact_spanish_edge_plans_return_expected_veterinary_subjects(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    repository = InMemorySearchRepository()
    vault_id = build_storage_namespace(
        network_kind="test", jurisdiction="ES", sector="onehealth-research", tenant_id="fixture",
    )
    for subject in _pipeline_subjects(generated["veterinary-study-a"]):
        repository.upsert(vault_id=vault_id, resource_type="ResearchSubject", resource=subject)
        for resource in subject["contained"]:
            repository.upsert(
                vault_id=vault_id,
                resource_type=str(resource["resourceType"]),
                resource=resource,
            )
    manager = ConversionSearchManager(SimpleNamespace(
        settings=SimpleNamespace(demo_mode=True, network_mode="test"),
        search_repo=repository,
    ))

    cases = (
        (
            [
                {"name": "Immunization.vaccine-code:text", "valueString": "rabia"},
                {"name": "Immunization.date", "valueString": "ge2025-01-01"},
            ],
            {VETERINARY_SUBJECTS[1]},
        ),
        (
            [
                {"name": "ResearchSubject.birthyear", "valueString": "lt2020"},
                {"name": "Condition.code:text", "valueString": "otitis"},
                {"name": "Procedure.code:text", "valueString": "radiografía"},
                {"name": "Procedure.date", "valueString": "ge2024-01-01"},
            ],
            {VETERINARY_SUBJECTS[0]},
        ),
    )
    for edge_parameters, expected in cases:
        result = manager.handle(
            tenant_id="fixture",
            jurisdiction="ES",
            sector="onehealth-research",
            resource_type="ResearchSubject",
            response=SimpleNamespace(),
            request=SimpleNamespace(headers={}, query_params={}),
            body={"resourceType": "Parameters", "parameter": [
                {"name": "ResearchSubject.study", "valueReference": {"reference": RESEARCH_STUDY_A}},
                *edge_parameters,
            ]},
        )
        actual = {
            entry["resource"]["meta"]["claims"]["ResearchSubject.identifier"]
            for entry in result["entry"]
        }
        assert actual == expected


def test_selected_research_subject_materializes_a_document_bundle_for_the_shared_health_viewer(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    repository = InMemorySearchRepository()
    vault_id = build_storage_namespace(
        network_kind="test", jurisdiction="ES", sector="onehealth-research", tenant_id="fixture",
    )
    subject = _pipeline_subjects(generated["veterinary-study-a"])[0]
    repository.upsert(vault_id=vault_id, resource_type="ResearchSubject", resource=subject)
    manager = ConversionSearchManager(SimpleNamespace(
        settings=SimpleNamespace(demo_mode=True, network_mode="test"),
        search_repo=repository,
    ))

    result = manager.handle_summary(
        tenant_id="fixture",
        jurisdiction="ES",
        sector="onehealth-research",
        request=SimpleNamespace(headers={}),
        body={"resourceType": "Parameters", "parameter": [
            {"name": "ResearchSubject.study", "valueReference": {"reference": RESEARCH_STUDY_A}},
            {"name": "ResearchSubject.identifier", "valueUri": VETERINARY_SUBJECTS[0]},
        ]},
    )

    assert result["resourceType"] == "Bundle"
    assert result["type"] == "document"
    assert result["entry"][0]["resource"]["resourceType"] == "Composition"
    section_references = {
        reference["reference"]
        for section in result["entry"][0]["resource"]["section"]
        for reference in section["entry"]
    }
    bundled_resources = {
        f'{entry["resource"]["resourceType"]}/{entry["resource"]["id"]}'
        for entry in result["entry"][1:]
    }
    assert section_references
    assert section_references <= bundled_resources


def test_researcher_saves_a_tagged_working_selection_without_mutating_the_canonical_twin(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    search = InMemorySearchRepository()
    vault = InMemoryVaultRepository()
    vault_id = build_storage_namespace(
        network_kind="test", jurisdiction="ES", sector="onehealth-research", tenant_id="fixture",
    )
    subject = _pipeline_subjects(generated["veterinary-study-a"])[0]
    search.upsert(vault_id=vault_id, resource_type="ResearchSubject", resource=subject)
    manager = ConversionSearchManager(SimpleNamespace(
        settings=SimpleNamespace(demo_mode=True, network_mode="test"),
        search_repo=search,
        vault_repo=vault,
    ))

    result = manager.handle_tag(
        tenant_id="fixture", jurisdiction="ES", sector="onehealth-research",
        request=SimpleNamespace(headers={}),
        body={"resourceType": "Parameters", "parameter": [
            {"name": "ResearchSubject.study", "valueReference": {"reference": RESEARCH_STUDY_A}},
            {"name": "ResearchSubject.identifier", "valueUri": VETERINARY_SUBJECTS[0]},
            {"name": "tag", "valueCoding": {"system": "urn:multibase:zResearcher", "code": "possible-candidate"}},
        ]},
    )

    assert result["resourceType"] == "Composition"
    claims = result["meta"]["claims"]
    assert claims["@type"] == "Composition:ResearcherWorkingSelection"
    assert claims["Composition.subject"] == VETERINARY_SUBJECTS[0]
    assert claims["Composition.meta-tag"] == "urn:multibase:zResearcher|possible-candidate"
    assert claims["Composition.userSelected"] == "true"
    assert search.search(vault_id=vault_id, resource_type="ResearchSubject", search_params={"identifier": VETERINARY_SUBJECTS[0]})[0] == subject

    removed = manager.handle_tag(
        tenant_id="fixture", jurisdiction="ES", sector="onehealth-research",
        request=SimpleNamespace(headers={}),
        body={"resourceType": "Parameters", "parameter": [
            {"name": "ResearchSubject.study", "valueReference": {"reference": RESEARCH_STUDY_A}},
            {"name": "ResearchSubject.identifier", "valueUri": VETERINARY_SUBJECTS[0]},
            {"name": "tag", "valueCoding": {"system": "urn:multibase:zResearcher", "code": "possible-candidate"}},
            {"name": "selected", "valueBoolean": False},
        ]},
    )

    assert removed == {"removed": True, "tag": "possible-candidate"}
    assert vault.query(vault_id, {
        "Composition.subject": VETERINARY_SUBJECTS[0],
        "Composition.meta-tag": "urn:multibase:zResearcher|possible-candidate",
    }, "Composition") == []


def test_source_event_dates_are_searchable_on_every_supported_clinical_family(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    subjects = _pipeline_subjects(generated["veterinary-study-a"])
    positive = subjects[0]
    resources = {resource["resourceType"]: resource for resource in positive["contained"]}

    expected_date_claims = {
        "AllergyIntolerance": "AllergyIntolerance.date",
        "Condition": "Condition.onset-datetime",
        "DiagnosticReport": "DiagnosticReport.date",
        "Immunization": "Immunization.date",
        "MedicationStatement": "MedicationStatement.effective",
        "Procedure": "Procedure.date",
    }
    for resource_type, claim in expected_date_claims.items():
        assert str(resources[resource_type]["meta"]["claims"].get(claim, "")).startswith(("2024-", "2025-"))


def test_committed_excel_fixtures_match_the_generator(tmp_path: Path) -> None:
    generated = generate_research_search_workbooks(tmp_path)
    fixture_dir = Path(__file__).parents[1] / "examples" / "research-search"
    for generated_path in generated.values():
        committed_path = fixture_dir / generated_path.name
        assert committed_path.is_file()
        assert _workbook_values(committed_path) == _workbook_values(generated_path)
