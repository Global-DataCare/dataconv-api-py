# Copyright Conéctate Soluciones y Aplicaciones SL
# SPDX-License-Identifier: Apache-2.0

"""Deterministic source-shaped Excel fixtures for research-search acceptance tests."""

from __future__ import annotations

from dataclasses import dataclass
import json
from pathlib import Path
from typing import Any


RESEARCH_STUDY_A = "ResearchStudy/provisional-search-a"
RESEARCH_STUDY_B = "ResearchStudy/provisional-search-b"

VETERINARY_ORIGINAL_PATIENT_IDS = tuple(
    f"SYNTHETIC-VET-PATIENT-{index:03d}" for index in range(1, 5)
)
HUMAN_ORIGINAL_PATIENT_IDS = tuple(
    f"SYNTHETIC-HUMAN-PATIENT-{index:03d}" for index in range(1, 5)
)
FIXTURE_SUBJECT_UUID_BY_ORIGINAL_ID = {
    **{
        original_id: f"{index}{index}{index}{index}{index}{index}{index}{index}-{index}{index}{index}{index}-4{index}{index}{index}-8{index}{index}{index}-{index}{index}{index}{index}{index}{index}{index}{index}{index}{index}{index}{index}"
        for index, original_id in enumerate(VETERINARY_ORIGINAL_PATIENT_IDS, start=1)
    },
    **{
        original_id: f"{index}{index}{index}{index}{index}{index}{index}{index}-{index}{index}{index}{index}-4{index}{index}{index}-8{index}{index}{index}-{index}{index}{index}{index}{index}{index}{index}{index}{index}{index}{index}{index}"
        for index, original_id in enumerate(HUMAN_ORIGINAL_PATIENT_IDS, start=5)
    },
}
VETERINARY_SUBJECTS = tuple(
    f"urn:uuid:{FIXTURE_SUBJECT_UUID_BY_ORIGINAL_ID[value]}"
    for value in VETERINARY_ORIGINAL_PATIENT_IDS
)
HUMAN_SUBJECTS = tuple(
    f"urn:uuid:{FIXTURE_SUBJECT_UUID_BY_ORIGINAL_ID[value]}"
    for value in HUMAN_ORIGINAL_PATIENT_IDS
)


@dataclass(frozen=True)
class FixtureRow:
    date: str
    original_patient_id: str
    birthyear: int
    species: str
    birthsex: str
    section: str
    family: str
    subfamily: str
    concept: str

    def values(self) -> tuple[Any, ...]:
        return (
            self.date,
            self.original_patient_id,
            self.birthyear,
            self.species,
            self.birthsex,
            self.section,
            self.family,
            self.subfamily,
            self.concept,
        )


DATA_FIELDS = (
    "date",
    "personal_id",
    "subject_birthyear",
    "subject_animal-species",
    "subject_birthsex",
    "section",
    "family",
    "subfamily",
    "concept",
)
DATA_HEADERS = (
    "FECHA",
    "ORIGINAL_PATIENT_ID",
    "AÑO_NACIMIENTO",
    "ESPECIE",
    "SEXO",
    "SECCION",
    "FAMILIA",
    "SUBFAMILIA",
    "CONCEPTO",
)
SEARCH_HEADERS = (
    "CASE_ID",
    "COMPLEXITY",
    "QUERY_TEXT_ES",
    "RESEARCH_STUDY",
    "PLAN_KIND",
    "FHIR_SEARCH_PLAN_JSON",
    "EXPECTED_RESEARCH_SUBJECT_IDS_JSON",
    "EXECUTION_STATUS",
    "NOTES",
)


def _row(
    original_patient_id: str,
    *,
    date: str,
    birthyear: int,
    species: str,
    birthsex: str,
    family: str,
    subfamily: str,
    concept: str,
) -> FixtureRow:
    return FixtureRow(
        date=date,
        original_patient_id=original_patient_id,
        birthyear=birthyear,
        species=species,
        birthsex=birthsex,
        section="clinica",
        family=family,
        subfamily=subfamily,
        concept=concept,
    )


def _veterinary_profile(original_id: str, *, birthyear: int) -> tuple[FixtureRow, ...]:
    common = {"original_patient_id": original_id, "birthyear": birthyear, "species": "CANINA", "birthsex": "Hembra"}
    return (
        _row(**common, date="2024-01-10", family="vacunas", subfamily="VACUNA PERRO", concept="VACUNA RABIA"),
        _row(**common, date="2025-02-15", family="diagnostico", subfamily="OTORRINO", concept="OTITIS EXTERNA"),
        _row(**common, date="2025-03-20", family="procedimiento", subfamily="RADIOLOGIA", concept="RADIOGRAFIA"),
        _row(**common, date="2025-04-10", family="alergias", subfamily="MEDICAMENTOS", concept="ALERGIA A PENICILINA"),
        _row(**common, date="2025-05-01", family="laboratorio", subfamily="ANALITICA", concept="HEMOGRAMA"),
        _row(**common, date="2025-05-05", family="medicamentos", subfamily="ANALGESIA", concept="GABAPENTINA 100 MG"),
        _row(**common, date="2025-06-01", family="consulta", subfamily="REVISION", concept="CONSULTA CLINICA"),
    )


def _veterinary_rows(study: str) -> tuple[FixtureRow, ...]:
    if study == RESEARCH_STUDY_B:
        return _veterinary_profile(VETERINARY_ORIGINAL_PATIENT_IDS[3], birthyear=2018)
    subject_b = VETERINARY_ORIGINAL_PATIENT_IDS[1]
    subject_c = VETERINARY_ORIGINAL_PATIENT_IDS[2]
    return (
        *_veterinary_profile(VETERINARY_ORIGINAL_PATIENT_IDS[0], birthyear=2018),
        _row(subject_b, date="2026-02-10", birthyear=2021, species="CANINA", birthsex="Macho", family="vacunas", subfamily="VACUNA PERRO", concept="VACUNA RABIA"),
        _row(subject_b, date="2026-03-03", birthyear=2021, species="CANINA", birthsex="Macho", family="diagnostico", subfamily="TRAUMATOLOGIA", concept="FRACTURA"),
        _row(subject_b, date="2026-03-04", birthyear=2021, species="CANINA", birthsex="Macho", family="procedimiento", subfamily="CIRUGIA", concept="CIRUGIA"),
        _row(subject_b, date="2026-03-05", birthyear=2021, species="CANINA", birthsex="Macho", family="laboratorio", subfamily="ANALITICA", concept="BIOQUIMICA"),
        _row(subject_b, date="2026-03-06", birthyear=2021, species="CANINA", birthsex="Macho", family="consulta", subfamily="URGENCIA", concept="CONSULTA DE URGENCIA"),
        _row(subject_c, date="2023-06-15", birthyear=2014, species="CANINA", birthsex="Hembra", family="vacunas", subfamily="VACUNA LEISH", concept="VACUNA LEISHMANIA"),
        _row(subject_c, date="2024-07-01", birthyear=2014, species="CANINA", birthsex="Hembra", family="diagnostico", subfamily="DERMATOLOGIA", concept="REACCION ALERGICA"),
        _row(subject_c, date="2024-07-02", birthyear=2014, species="CANINA", birthsex="Hembra", family="alergias", subfamily="AMBIENTAL", concept="ALERGIA AMBIENTAL"),
        _row(subject_c, date="2024-08-01", birthyear=2014, species="CANINA", birthsex="Hembra", family="procedimiento", subfamily="ECOGRAFIA", concept="ECOGRAFIA ABDOMINAL"),
        _row(subject_c, date="2024-08-02", birthyear=2014, species="CANINA", birthsex="Hembra", family="medicamentos", subfamily="ANTIINFLAMATORIO", concept="MELOXICAM 1.5 MG"),
    )


def _human_positive_profile(original_id: str) -> tuple[FixtureRow, ...]:
    common = {"original_patient_id": original_id, "birthyear": 1980, "species": "", "birthsex": "female"}
    return (
        _row(**common, date="2025-10-15", family="vacunas", subfamily="INFLUENZA", concept="SE ADMINISTRA VACUNA GRIPE"),
        _row(**common, date="2025-11-02", family="diagnostico", subfamily="ENDOCRINOLOGIA", concept="DIABETES MELLITUS"),
        _row(**common, date="2025-11-03", family="laboratorio", subfamily="HBA1C", concept="HbA1c 8,4 %"),
        _row(**common, date="2025-11-04", family="medicamentos", subfamily="ANTIDIABETICO", concept="METFORMINA 850 MG"),
        _row(**common, date="2025-11-05", family="procedimiento", subfamily="RADIOLOGIA", concept="RADIOLOGIA"),
        _row(**common, date="2025-11-06", family="alergias", subfamily="MEDICAMENTOS", concept="ALERGIA A PENICILINA"),
    )


def _human_rows(study: str) -> tuple[FixtureRow, ...]:
    if study == RESEARCH_STUDY_B:
        return _human_positive_profile(HUMAN_ORIGINAL_PATIENT_IDS[3])
    subject_b = HUMAN_ORIGINAL_PATIENT_IDS[1]
    subject_c = HUMAN_ORIGINAL_PATIENT_IDS[2]
    return (
        *_human_positive_profile(HUMAN_ORIGINAL_PATIENT_IDS[0]),
        _row(subject_b, date="2025-10-15", birthyear=2018, species="", birthsex="male", family="vacunas", subfamily="INFLUENZA", concept="SE ADMINISTRA VACUNA GRIPE"),
        _row(subject_b, date="2025-11-02", birthyear=2018, species="", birthsex="male", family="diagnostico", subfamily="ENDOCRINOLOGIA", concept="DIABETES MELLITUS"),
        _row(subject_b, date="2025-11-03", birthyear=2018, species="", birthsex="male", family="laboratorio", subfamily="HBA1C", concept="HbA1c 6,2 %"),
        _row(subject_b, date="2025-11-04", birthyear=2018, species="", birthsex="male", family="medicamentos", subfamily="ANTIDIABETICO", concept="METFORMINA 850 MG"),
        _row(subject_c, date="2023-01-05", birthyear=1955, species="", birthsex="female", family="laboratorio", subfamily="HBA1C", concept="HbA1c 9,1 %"),
        _row(subject_c, date="2026-01-20", birthyear=1955, species="", birthsex="female", family="diagnostico", subfamily="CARDIOLOGIA", concept="HIPERTENSION ARTERIAL"),
        _row(subject_c, date="2026-01-23", birthyear=1955, species="", birthsex="female", family="medicamentos", subfamily="ANTIHIPERTENSIVO", concept="ENALAPRIL 10 MG"),
    )


def _plan(resource_type: str, params: dict[str, Any], **extra: Any) -> dict[str, Any]:
    return {"resourceType": resource_type, "params": params, **extra}


def _search_cases(cohort: str, study: str) -> tuple[tuple[Any, ...], ...]:
    study_b = study == RESEARCH_STUDY_B
    subjects = HUMAN_SUBJECTS if cohort == "human" else VETERINARY_SUBJECTS
    expected_profile = (subjects[3],) if study_b else (subjects[0],)
    if cohort == "human":
        cases = (
            ("HUM-S-GRIPE", "simple", "Vacunados de gripe", [_plan("Immunization", {"vaccine-code:text": "gripe"})], (subjects[3],) if study_b else subjects[:2], "single-resource", "executable", "Vacunación y fecha se derivan de la fila fuente."),
            ("HUM-S-DIABETES", "simple", "Diagnosticados de diabetes", [_plan("Condition", {"code:text": "diabetes"})], (subjects[3],) if study_b else subjects[:2], "single-resource", "executable", "Texto local pendiente de confirmación terminológica."),
            ("HUM-C-DIABETES-HBA1C", "complex", "Diabetes con HbA1c mayor de 8 desde 2025 y metformina", [_plan("Condition", {"code:text": "diabetes"}), _plan("Observation", {"code:text": "hba1c", "date": "ge2025-01-01", "value-quantity": "gt8"}), _plan("MedicationStatement", {"code:text": "metformina"})], expected_profile, "subject-intersection", "planner-required", "A es positivo; B falla por valor; C falla por condición o fecha; D queda aislado en Study B."),
            ("HUM-C-BIRTH-GRIPE", "complex", "Nacidos antes de 2000 con vacuna de gripe desde 2025", [_plan("ResearchSubject", {"study": study}, claimFilters={"Subject.birthyear": "lt2000"}), _plan("Immunization", {"vaccine-code:text": "gripe", "date": "ge2025-01-01"})], expected_profile, "subject-intersection", "planner-required", "Cruce de demografía y vacunación."),
        )
    else:
        cases = (
            ("VET-S-RABIA", "simple", "Vacunados de rabia", [_plan("Immunization", {"vaccine-code:text": "rabia"})], (subjects[3],) if study_b else subjects[:2], "single-resource", "executable", "Vacunación derivada del concepto."),
            ("VET-C-RABIA-DATE", "complex", "Vacuna de rabia desde 2025", [_plan("Immunization", {"vaccine-code:text": "rabia", "date": "ge2025-01-01"})], () if study_b else (subjects[1],), "single-resource", "executable", "Filtro por fecha de vacunación."),
            ("VET-C-OTITIS-RADIO", "complex", "Otitis y radiografía", [_plan("Condition", {"code:text": "otitis"}), _plan("Procedure", {"code:text": "radiografia"})], expected_profile, "subject-intersection", "planner-required", "AND entre recursos del mismo sujeto."),
            ("VET-C-ALLERGY-REPORT", "complex", "Alergia y hemograma", [_plan("AllergyIntolerance", {"code:text": "alergia"}), _plan("DiagnosticReport", {"code:text": "hemograma"})], expected_profile, "subject-intersection", "planner-required", "Cruce de alergia e informe diagnóstico."),
        )
    return tuple(
        (
            case_id,
            complexity,
            query_text,
            study,
            plan_kind,
            json.dumps(plan, ensure_ascii=False, separators=(",", ":")),
            json.dumps(list(expected), separators=(",", ":")),
            status,
            notes,
        )
        for case_id, complexity, query_text, plan, expected, plan_kind, status, notes in cases
    )


def _write_workbook(
    path: Path,
    *,
    cohort: str,
    study: str,
    rows: tuple[FixtureRow, ...],
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - optional dependency guard
        raise RuntimeError("Install the Excel extra before generating research fixtures") from exc

    workbook = Workbook()
    data = workbook.active
    data.title = "DATA"
    subject_kind = "animal" if cohort == "veterinary" else "person"
    study_slug = "a" if study == RESEARCH_STUDY_A else "b"
    data.append([
        "API-CONFIG:language=es:software-id=research-search-fixture-"
        f"{cohort}-study-{study_slug}:subjectKind={subject_kind}:dataUse=secondary"
    ])
    data.append(DATA_FIELDS)
    data.append(DATA_HEADERS)
    for row in rows:
        data.append(row.values())

    searches = workbook.create_sheet("SEARCH_CASES")
    searches.append(SEARCH_HEADERS)
    for row in _search_cases(cohort, study):
        searches.append(row)

    readme = workbook.create_sheet("README")
    readme.append(("CAMPO", "VALOR"))
    readme.append(("RESEARCH_STUDY", study))
    readme.append(("COHORTE", cohort))
    readme.append(("PRIVACIDAD", "Todos los identificadores y datos demográficos son sintéticos."))
    readme.append(("IDENTIDAD", "ORIGINAL_PATIENT_ID se resuelve confidencialmente a un UUID estable; no viaja un UUID en DATA."))
    readme.append(("ORIGEN_CONCEPTOS", "Textos representativos tomados de patrones de las tablas suministradas."))
    readme.append(("USO", "Fixture reproducible para importación, revisión y planificación de búsquedas."))
    readme.append(("LIMITACION", "subject-intersection requiere planificador y correlación por sujeto."))

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet, header_row in ((data, 3), (searches, 1), (readme, 1)):
        for cell in sheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(row, column).value or "") for row in range(1, sheet.max_row + 1)]
            sheet.column_dimensions[get_column_letter(column)].width = min(max(map(len, values)) + 2, 48)

    workbook.properties.title = f"Research search {cohort} {study_slug} fixture"
    workbook.properties.subject = "Anonymous clinical research search acceptance data"
    workbook.properties.creator = "DataConv"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def generate_research_search_workbooks(output_dir: Path) -> dict[str, Path]:
    """Generate source-shaped cohort files for both provisional studies."""

    output_dir = Path(output_dir)
    specifications = (
        ("veterinary-study-a", "veterinary", RESEARCH_STUDY_A, _veterinary_rows(RESEARCH_STUDY_A)),
        ("veterinary-study-b", "veterinary", RESEARCH_STUDY_B, _veterinary_rows(RESEARCH_STUDY_B)),
        ("human-study-a", "human", RESEARCH_STUDY_A, _human_rows(RESEARCH_STUDY_A)),
        ("human-study-b", "human", RESEARCH_STUDY_B, _human_rows(RESEARCH_STUDY_B)),
    )
    generated: dict[str, Path] = {}
    for key, cohort, study, rows in specifications:
        path = output_dir / f"research-search-{key}.xlsx"
        _write_workbook(path, cohort=cohort, study=study, rows=rows)
        generated[key] = path
    return generated
