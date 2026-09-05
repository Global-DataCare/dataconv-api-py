# Flow contract: classify legacy ingestion fields explicitly; never infer canonical claims from physical-looking names.

from adapter_ingestion.base_config_contract import (
    BASE_CONFIG_FIELDS,
    BaseConfigFieldKind,
    canonical_claim_for_base_config_field,
    reconcile_base_config_workbook,
)
from adapter_ingestion.claim_governance import (
    DATACONV_CLAIM_ALIASES,
    DATACONV_EXTENSION_CLAIMS,
)
from openpyxl import Workbook, load_workbook


def test_all_legacy_dataconv_claims_are_classified_as_aliases_or_extensions() -> None:
    expected = {
        "DocumentReference.text",
        "Encounter.date",
        "Encounter.servicetype",
        "OperationOutcome.issueCode",
        "OperationOutcome.rowNumber",
        "OperationOutcome.sectionFamily",
        "ResearchSubject.identifier",
        "ResearchSubject.status",
        "Subject.active",
        "Subject.animal-breed",
        "Subject.animal-genderstatus",
        "Subject.animal-species",
        "Subject.birthsex",
        "Subject.birthyear",
        "Subject.id",
        "Subject.language",
        "Subject.link",
    }
    assert set(DATACONV_CLAIM_ALIASES) | DATACONV_EXTENSION_CLAIMS == expected
    assert set(DATACONV_CLAIM_ALIASES).isdisjoint(DATACONV_EXTENSION_CLAIMS)
    assert DATACONV_CLAIM_ALIASES == {
        "Encounter.date": "Encounter.period-start",
        "Encounter.servicetype": "Encounter.type",
    }


def test_base_config_separates_control_canonical_extension_and_pending_fields() -> None:
    assert len(BASE_CONFIG_FIELDS) == 50
    assert {entry.kind for entry in BASE_CONFIG_FIELDS.values()} == {
        BaseConfigFieldKind.CONTROL,
        BaseConfigFieldKind.CANONICAL_CLAIM,
        BaseConfigFieldKind.DATACONV_EXTENSION,
        BaseConfigFieldKind.PENDING,
    }
    assert canonical_claim_for_base_config_field("invoice_date") == "Invoice.date"
    assert canonical_claim_for_base_config_field("chargeitem_productcode") == (
        "ChargeItem.code"
    )
    assert canonical_claim_for_base_config_field("chargeitem_productname") == (
        "ChargeItem.code-text"
    )
    assert canonical_claim_for_base_config_field("chargeitem_quantity") == (
        "ChargeItem.quantity-number"
    )
    assert canonical_claim_for_base_config_field("chargeitem_unit") == (
        "ChargeItem.quantity-unit"
    )
    assert BASE_CONFIG_FIELDS["origin"].kind is BaseConfigFieldKind.PENDING


def test_base_config_workbook_is_split_without_overwriting_the_source(tmp_path) -> None:
    source = tmp_path / "mapping.xlsx"
    output = tmp_path / "mapping-reconciled.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BaseConfig"
    sheet.append(("API config.", "Data", "Local", "Description"))
    sheet.append(("section", "string", "SECCION", "Control"))
    sheet.append(("invoice_date", "date", "FECHA_FACTURA", "Invoice date"))
    sheet.append(("origin", "string", "ORIGEN", "Ungoverned"))
    workbook.save(source)

    reconcile_base_config_workbook(source, output)

    assert source.exists()
    reconciled = load_workbook(output, read_only=True, data_only=True)
    assert {
        "BaseConfig-Control",
        "BaseConfig-Canonical",
        "BaseConfig-Extensions",
        "BaseConfig-Pending",
    }.issubset(reconciled.sheetnames)
    canonical_rows = list(reconciled["BaseConfig-Canonical"].iter_rows(values_only=True))
    assert canonical_rows[1][0] == "invoice_date"
    assert canonical_rows[1][4] == "Invoice.date"
