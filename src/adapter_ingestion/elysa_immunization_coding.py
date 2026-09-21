"""Reviewed ATC/ATCvet candidate enrichment for the Elysa API-CONFIG workbook."""

# Copyright ConnectHealth.info (Connecting Solution & Applications Ltd.)
# SPDX-License-Identifier: Apache-2.0

from __future__ import annotations

from collections import OrderedDict
import csv
from dataclasses import dataclass, field
from io import StringIO
from pathlib import Path
import re
import unicodedata


ATC_SYSTEM = "http://www.whocc.no/atc"
ATCVET_SYSTEM = "http://www.whocc.no/atcvet"

VACCINE_CODE_CLAIM = "Immunization.vaccine-code"
VACCINE_DISPLAY_CLAIM = "Immunization.vaccine-code-display"
VACCINE_TEXT_CLAIM = "Immunization.vaccine-code-text"


@dataclass(frozen=True)
class Coding:
    system: str
    code: str
    display: str
    text_es: str


@dataclass(frozen=True)
class Match:
    codings: tuple[Coding, ...]
    species_codes: str
    species_display: str
    confidence: str
    basis: str


@dataclass
class SummaryItem:
    sheet: str
    source_column: str
    original_text: str
    codes: str
    displays: str
    texts_es: str
    species_codes: str
    species_display: str
    confidence: str
    basis: str
    row_numbers: list[int] = field(default_factory=list)

    @property
    def matched_rows(self) -> int:
        return len(self.row_numbers)


@dataclass(frozen=True)
class EnrichmentSummary:
    items: tuple[SummaryItem, ...]
    matched_rows: int
    matched_rows_by_sheet: dict[str, int]


def _coding(code: str, display: str, text_es: str, *, human: bool = False) -> Coding:
    return Coding(ATC_SYSTEM if human else ATCVET_SYSTEM, code, display, text_es)


CODINGS = {
    "QI07AA02": _coding("QI07AA02", "rabies virus", "virus de la rabia"),
    "QI07AD04": _coding(
        "QI07AD04",
        "canine distemper virus + canine adenovirus + canine parvovirus + canine parainfluenza virus",
        "virus del moquillo canino + adenovirus canino + parvovirus canino + virus de la parainfluenza canina",
    ),
    "QI07AB01": _coding("QI07AB01", "leptospira", "leptospira"),
    "QI07AO01": _coding("QI07AO01", "leishmania", "leishmania"),
    "QI07AI02": _coding(
        "QI07AI02",
        "live canine distemper virus + live canine adenovirus + live canine parainfluenza virus + live canine parvovirus + inactivated leptospira",
        "virus vivo del moquillo canino + adenovirus canino vivo + virus vivo de la parainfluenza canina + parvovirus canino vivo + leptospira inactivada",
    ),
    "QI07AJ06": _coding(
        "QI07AJ06",
        "live canine distemper virus + live canine adenovirus + live parainfluenza virus + live canine parvovirus + inactivated rabies + inactivated leptospira",
        "virus vivo del moquillo canino + adenovirus canino vivo + virus vivo de la parainfluenza canina + parvovirus canino vivo + rabia inactivada + leptospira inactivada",
    ),
    "QI06AD04": _coding(
        "QI06AD04",
        "feline panleucopenia virus / parvovirus + feline rhinotracheitis virus + feline calicivirus",
        "virus de la panleucopenia felina / parvovirus + virus de la rinotraqueitis felina + calicivirus felino",
    ),
    "QI06AH09": _coding(
        "QI06AH09",
        "live feline rhinotracheitis virus + live feline panleucopenia virus / parvovirus + inactivated feline calicivirus antigen",
        "virus vivo de la rinotraqueitis felina + virus vivo de la panleucopenia felina / parvovirus + antigeno inactivado de calicivirus felino",
    ),
    "QI07AD02": _coding(
        "QI07AD02",
        "canine distemper virus + canine adenovirus + canine parvovirus",
        "virus del moquillo canino + adenovirus canino + parvovirus canino",
    ),
    "QI06AA01": _coding("QI06AA01", "feline leukaemia virus", "virus de la leucemia felina"),
    "QI06AH10": _coding(
        "QI06AH10",
        "live feline rhinotracheitis virus + live feline panleucopenia virus / parvovirus + inactivated feline calicivirus + feline leukaemia, recombinant live canarypox virus",
        "virus vivo de la rinotraqueitis felina + virus vivo de la panleucopenia felina / parvovirus + calicivirus felino inactivado + leucemia felina, virus vivo recombinante de la viruela del canario",
    ),
    "QI07AE01": _coding("QI07AE01", "bordetella", "bordetella"),
    "QI07AD01": _coding("QI07AD01", "canine parvovirus", "parvovirus canino"),
    "QI06AD07": _coding(
        "QI06AD07",
        "feline leukaemia, recombinant live canarypox virus",
        "leucemia felina, virus vivo recombinante de la viruela del canario",
    ),
    "QI07AD03": _coding(
        "QI07AD03",
        "canine distemper virus + canine parvovirus",
        "virus del moquillo canino + parvovirus canino",
    ),
    "QI07AF01": _coding(
        "QI07AF01",
        "bordetella + canine parainfluenza virus",
        "bordetella + virus de la parainfluenza canina",
    ),
    "QI06AJ03": _coding(
        "QI06AJ03",
        "live feline rhinotracheitis virus + inactivated feline calicivirus antigen + live feline panleucopenia virus / parvovirus + live chlamydia",
        "virus vivo de la rinotraqueitis felina + antigeno inactivado de calicivirus felino + virus vivo de la panleucopenia felina / parvovirus + clamidia viva",
    ),
    "QI07AL05": _coding(
        "QI07AL05",
        "bordetella + canine parainfluenza virus",
        "bordetella + virus de la parainfluenza canina",
    ),
    "QI06AH07": _coding(
        "QI06AH07",
        "live feline panleucopenia virus / parvovirus + live feline rhinotracheitis virus + live feline calicivirus + inactivated feline leukaemia virus",
        "virus vivo de la panleucopenia felina / parvovirus + virus vivo de la rinotraqueitis felina + calicivirus felino vivo + virus inactivado de la leucemia felina",
    ),
    "QI06AJ05": _coding(
        "QI06AJ05",
        "live feline rhinotracheitis virus + inactivated feline calicivirus antigen + live feline panleucopenia virus / parvovirus + live chlamydia + feline leukaemia recombinant live canarypox virus",
        "virus vivo de la rinotraqueitis felina + antigeno inactivado de calicivirus felino + virus vivo de la panleucopenia felina / parvovirus + clamidia viva + leucemia felina con virus vivo recombinante de la viruela del canario",
    ),
    "QI08AH01": _coding(
        "QI08AH01",
        "live myxomatosis virus + inactivated rabbit haemorrhagic disease virus",
        "virus vivo de la mixomatosis + virus inactivado de la enfermedad hemorragica del conejo",
    ),
    "QI08AD": _coding("QI08AD", "Live viral vaccines", "Vacunas virales vivas"),
    "J07BB": _coding("J07BB", "Influenza vaccines", "Vacunas contra la gripe", human=True),
}


CANITAS_CODES = {
    "EURICAN R": ("QI07AA02",),
    "EURICAN DAPPI": ("QI07AD04",),
    "EURICAN L4": ("QI07AB01",),
    "LETIFEND": ("QI07AO01",),
    "VERSICAN PLUS DHPPI/L4": ("QI07AI02",),
    "VERSICAN PLUS DHPPI/L4R": ("QI07AJ06",),
    "VERSIFEL CVR": ("QI06AD04",),
    "VERSIGUARD RABIA": ("QI07AA02",),
    "PUREVAX RCP": ("QI06AH09",),
    "EURICAN DAP": ("QI07AD02",),
    "LEUCOGEN": ("QI06AA01",),
    "PUREVAX RCP + FELV": ("QI06AH10",),
    "VERSICAN PLUS BB ORAL": ("QI07AE01",),
    "VANGUARD-CPV": ("QI07AD01",),
    "PUREVAX FELV": ("QI06AD07",),
    "VERSIFEL FELV": ("QI06AA01",),
    "VERSICAN PLUS DP": ("QI07AD03",),
    "CANIGEN DHPPI/L": ("QI07AI02",),
    "NOBIVAC KC": ("QI07AF01",),
    "NOBIVAC RABIA": ("QI07AA02",),
    "VERSICAN PLUS BBPI IN": ("QI07AF01",),
    "RABIGEN-L": ("QI07AA02",),
    "VERSICAN PLUS DHP": ("QI07AD02",),
    "PUREVAX RCP-CH": ("QI06AJ03",),
    "EURICAN PNEUMO": ("QI07AL05",),
    "LEUCOFELIGEN CRP/FELV": ("QI06AH07",),
    "PUREVAX RCP-CH + FELV": ("QI06AJ05",),
    "FELIGEN CRP": ("QI06AD04",),
    "VANGUARD-7": ("QI07AI02",),
    "NOBIVAC MYXO-RHD PLUS": ("QI08AD",),
    "EURICAN PRIMO": ("QI07AD01",),
}

CAEIRA_CODES = {
    "VACUNA MIXOMATOSIS+ VIR. HEMORRAGICA": ("QI08AH01",),
    "VACUNA LEUCEMIA": ("QI06AA01", "QI06AD07"),
    "VACUNA TRIV. GATO": ("QI06AD04",),
    "VACUNA TRIVA . + LEUC.": ("QI06AH07",),
    "VACUNA EURICAN PNEUMO": ("QI07AL05",),
    "VACUNA LEPTO": ("QI07AB01",),
    "VACUNA NOVIBAC KC": ("QI07AF01",),
    "VACUNA RABIA": ("QI07AA02",),
    "VACUNA TETRA + LEPTO": ("QI07AI02",),
    "VACUNA TETRAVALENTE": ("QI07AD04",),
    "VERSICAN PLUS BB ORAL": ("QI07AE01",),
}


def _normalized(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(char for char in text if not unicodedata.combining(char)).upper().split())


def _species_for_codes(codes: tuple[str, ...]) -> tuple[str, str]:
    species: OrderedDict[str, str] = OrderedDict()
    for code in codes:
        if code.startswith("QI07"):
            species["9615"] = "dog"
        elif code.startswith("QI06"):
            species["9685"] = "cat"
        elif code.startswith("QI08"):
            species["9986"] = "rabbit"
        elif code.startswith("J07"):
            species["9606"] = "human"
    return ",".join(species), ",".join(species.values())


def _match(sheet: str, original: object, family: object, subfamily: object) -> Match | None:
    text = _normalized(original)
    if not text:
        return None

    codes: tuple[str, ...] | None = None
    confidence = "high"
    basis = "explicit vaccine product or administration text"

    if sheet == "Canitas 2" and _normalized(subfamily) == "VACUNAS":
        if text.startswith("EURICAN SOLVENTE"):
            return None
        product_name = next(
            (name for name in sorted(CANITAS_CODES, key=len, reverse=True) if text.startswith(name)),
            None,
        )
        codes = CANITAS_CODES.get(product_name) if product_name else None
    elif sheet == "CV A Caeira" and _normalized(family) == "VACUNAS":
        product_name = next(
            (name for name in sorted(CAEIRA_CODES, key=len, reverse=True) if text.startswith(name)),
            None,
        )
        codes = CAEIRA_CODES.get(product_name) if product_name else None
        if text == "VACUNA LEUCEMIA":
            confidence = "ambiguous"
            basis = "vaccine type is present but formulation is absent; two ATCvet candidates retained"
        elif text == "VACUNA RABIA":
            confidence = "medium"
            basis = "ATCvet main-species convention; actual subject species is absent from this sheet"
    elif sheet == "Pinol Vepahi":
        if "PENDIENTE" in text or "RECOM" in text or "REACCION" in text:
            return None
        found: list[str] = []
        patterns = (
            (r"EURICAN\s+R(?:\s|\d|$)", "QI07AA02"),
            (r"EURICAN\s+(?:DAPPI|DJHPPI)", "QI07AD04"),
            (r"EURICAN\s+L4", "QI07AB01"),
            (r"VACUNA\s+RABIA.*NOBIVAC", "QI07AA02"),
            (r"VACUNADA\s+DE\s+DHP", "QI07AD02"),
        )
        for pattern, code in patterns:
            if re.search(pattern, text) and code not in found:
                found.append(code)
        codes = tuple(found) or None
        if codes and "VACUNADA DE DHP" in text:
            confidence = "medium"
            basis = "historical administration in clinical narrative"
    elif sheet == "Survet Diagonal":
        if "VACUNACIONES DE LEISHMANIA" in text:
            codes = ("QI07AO01",)
            confidence = "medium"
            basis = "historical vaccination in clinical narrative"
        elif "VACUNADA EN JUNIO CON FELIGEN CRP" in text:
            codes = ("QI06AD04",)
            confidence = "high"
            basis = "historical product administration in clinical narrative"
        elif "LA VACUNAN CON LEUCOFELIGEN" in text:
            codes = ("QI06AH07",)
            confidence = "high"
            basis = "historical product administration in clinical narrative"
    elif sheet == "Sanios" and "SE ADMINISTRA VACUNA GRIPE" in text:
        codes = ("J07BB",)
        confidence = "medium"
        basis = "administration is explicit but vaccine formulation is absent; broad WHO ATC group retained"

    if not codes:
        return None
    species_codes, species_display = _species_for_codes(codes)
    return Match(
        codings=tuple(CODINGS[code] for code in codes),
        species_codes=species_codes,
        species_display=species_display,
        confidence=confidence,
        basis=basis,
    )


def _joined(match: Match) -> tuple[str, str, str]:
    def csv_value(values: tuple[str, ...]) -> str:
        output = StringIO()
        csv.writer(output, lineterminator="").writerow(values)
        return output.getvalue()

    return (
        csv_value(tuple(f"{coding.system}|{coding.code}" for coding in match.codings)),
        csv_value(tuple(coding.display for coding in match.codings)),
        csv_value(tuple(coding.text_es for coding in match.codings)),
    )


def enrich_workbook(source: Path | str, target: Path | str) -> EnrichmentSummary:
    """Create an enriched copy and return aggregated review evidence."""
    from openpyxl import load_workbook

    source_path = Path(source)
    target_path = Path(target)
    workbook = load_workbook(source_path)
    summary_by_key: OrderedDict[tuple[str, str], SummaryItem] = OrderedDict()
    matched_rows_by_sheet: dict[str, int] = {}

    for sheet in workbook.worksheets:
        mappings = [sheet.cell(2, column).value for column in range(1, sheet.max_column + 1)]
        claim_columns: list[int] = []
        for claim in (VACCINE_CODE_CLAIM, VACCINE_DISPLAY_CLAIM, VACCINE_TEXT_CLAIM):
            try:
                column = mappings.index(claim) + 1
            except ValueError:
                column = sheet.max_column + 1
                sheet.cell(2, column, claim)
                mappings.append(claim)
            claim_columns.append(column)
        for column, header in zip(
            claim_columns,
            ("IMMUNIZATION_VACCINE_CODE", "IMMUNIZATION_VACCINE_CODE_DISPLAY", "IMMUNIZATION_VACCINE_CODE_TEXT"),
        ):
            sheet.cell(3, column, header)

        mapping_to_column = {
            str(sheet.cell(2, column).value): column
            for column in range(1, sheet.max_column + 1)
            if sheet.cell(2, column).value
        }
        source_claim = next(
            (
                claim
                for claim in (
                    "concept",
                    "Condition.code-text",
                    "DiagnosticReport.code-text",
                )
                if claim in mapping_to_column
            ),
            "",
        )
        source_column = mapping_to_column.get(source_claim)
        family_column = mapping_to_column.get("family")
        subfamily_column = mapping_to_column.get("subfamily")
        matched_rows_by_sheet[sheet.title] = 0

        if source_column is None:
            continue
        for row_number in range(4, sheet.max_row + 1):
            original = sheet.cell(row_number, source_column).value
            match = _match(
                sheet.title,
                original,
                sheet.cell(row_number, family_column).value if family_column else None,
                sheet.cell(row_number, subfamily_column).value if subfamily_column else None,
            )
            if match is None:
                for column in claim_columns:
                    sheet.cell(row_number, column, None)
                continue
            code_value, display_value, text_value = _joined(match)
            for column, value in zip(claim_columns, (code_value, display_value, text_value)):
                sheet.cell(row_number, column, value)
            matched_rows_by_sheet[sheet.title] += 1

            key = (sheet.title, str(original))
            if key not in summary_by_key:
                summary_by_key[key] = SummaryItem(
                    sheet=sheet.title,
                    source_column=source_claim,
                    original_text=str(original),
                    codes=code_value,
                    displays=display_value,
                    texts_es=text_value,
                    species_codes=match.species_codes,
                    species_display=match.species_display,
                    confidence=match.confidence,
                    basis=match.basis,
                )
            summary_by_key[key].row_numbers.append(row_number)

    target_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target_path)
    workbook.close()
    items = tuple(summary_by_key.values())
    return EnrichmentSummary(
        items=items,
        matched_rows=sum(item.matched_rows for item in items),
        matched_rows_by_sheet=matched_rows_by_sheet,
    )
